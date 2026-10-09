from __future__ import annotations

import ast
import calendar
import json
import re
import time
from datetime import date, datetime, timedelta
from io import StringIO
from pathlib import Path
from typing import Any, cast

import httpx
import pandas as pd
import requests
import urllib3
from bs4 import BeautifulSoup
from openai import OpenAI
from openpyxl import load_workbook
from dateutil.relativedelta import relativedelta

# ============================================================
# NBA TRADING ENGINE — ARCHITECTURE
# Stage 1: Collect NBA contracts, transactions and other assets
# Stage 2: GPT evaluates NBA players
# Stage 3: GPT interprets the user's trade request
# Stage 4: GPT separates required/optional acquisition and outgoing pools
# Stage 5: Individual CBA eligibility checks (no salary calculation)
# Stage 6: Trade generation + salary feasibility — DISABLED
# Stage 7: Full multi-party CBA validation — DISABLED
# Stage 8: Legal trade results — DISABLED
# ============================================================


def big_boy(user_statement: str, api_key: str):
    """Run the pipeline through player eligibility; print four final dictionaries.

    Both arguments must be strings. All helper functions, data, client objects,
    and intermediate values live inside this function.
    """
    if not isinstance(user_statement, str) or not user_statement.strip():
        raise ValueError("user_statement must be a nonempty string")
    if not isinstance(api_key, str) or not api_key.strip():
        raise ValueError("api_key must be a nonempty string")

    # ============================================================
    # STAGE 1 — NBA DATA COLLECTION: CONFIGURATION AND HELPERS
    # ============================================================
    SPOTRAC_TEAMS = {

        "atlanta-hawks": ("Atlanta Hawks", "ATL"), "boston-celtics": ("Boston Celtics", "BOS"), "brooklyn-nets": ("Brooklyn Nets", "BKN"),

        "charlotte-hornets": ("Charlotte Hornets", "CHA"), "chicago-bulls": ("Chicago Bulls", "CHI"), "cleveland-cavaliers": ("Cleveland Cavaliers", "CLE"),

        "dallas-mavericks": ("Dallas Mavericks", "DAL"), "denver-nuggets": ("Denver Nuggets", "DEN"), "detroit-pistons": ("Detroit Pistons", "DET"),

        "golden-state-warriors": ("Golden State Warriors", "GSW"), "houston-rockets": ("Houston Rockets", "HOU"), "indiana-pacers": ("Indiana Pacers", "IND"),

        "la-clippers": ("Los Angeles Clippers", "LAC"), "los-angeles-lakers": ("Los Angeles Lakers", "LAL"), "memphis-grizzlies": ("Memphis Grizzlies", "MEM"),

        "miami-heat": ("Miami Heat", "MIA"), "milwaukee-bucks": ("Milwaukee Bucks", "MIL"), "minnesota-timberwolves": ("Minnesota Timberwolves", "MIN"),

        "new-orleans-pelicans": ("New Orleans Pelicans", "NOP"), "new-york-knicks": ("New York Knicks", "NYK"), "oklahoma-city-thunder": ("Oklahoma City Thunder", "OKC"),

        "orlando-magic": ("Orlando Magic", "ORL"), "philadelphia-76ers": ("Philadelphia 76ers", "PHI"), "phoenix-suns": ("Phoenix Suns", "PHX"),

        "portland-trail-blazers": ("Portland Trail Blazers", "POR"), "sacramento-kings": ("Sacramento Kings", "SAC"), "san-antonio-spurs": ("San Antonio Spurs", "SAS"),

        "toronto-raptors": ("Toronto Raptors", "TOR"), "utah-jazz": ("Utah Jazz", "UTA"), "washington-wizards": ("Washington Wizards", "WAS")

    }

    PICK_ALIASES = {

        "Atlanta Hawks": ["Atlanta"], "Boston Celtics": ["Boston"], "Brooklyn Nets": ["Brooklyn"], "Charlotte Hornets": ["Charlotte"],

        "Chicago Bulls": ["Chicago"], "Cleveland Cavaliers": ["Cleveland"], "Dallas Mavericks": ["Dallas"], "Denver Nuggets": ["Denver"],

        "Detroit Pistons": ["Detroit"], "Golden State Warriors": ["Golden State"], "Houston Rockets": ["Houston"], "Indiana Pacers": ["Indiana"],

        "Los Angeles Clippers": ["L.A. Clippers", "LA Clippers", "Los Angeles Clippers"], "Los Angeles Lakers": ["L.A. Lakers", "LA Lakers", "Los Angeles Lakers"],

        "Memphis Grizzlies": ["Memphis"], "Miami Heat": ["Miami"], "Milwaukee Bucks": ["Milwaukee"], "Minnesota Timberwolves": ["Minnesota"],

        "New Orleans Pelicans": ["New Orleans"], "New York Knicks": ["New York"], "Oklahoma City Thunder": ["Oklahoma City"], "Orlando Magic": ["Orlando"],

        "Philadelphia 76ers": ["Philadelphia"], "Phoenix Suns": ["Phoenix"], "Portland Trail Blazers": ["Portland"], "Sacramento Kings": ["Sacramento"],

        "San Antonio Spurs": ["San Antonio"], "Toronto Raptors": ["Toronto"], "Utah Jazz": ["Utah"], "Washington Wizards": ["Washington"]

    }

    MIN_CONTRACT_SEASON = 2026

    TRANSACTION_START = "2021-07-01"

    REQUEST_DELAY = 0.25

    SAVE_EXCEL = True

    PLAYER_POSITIONS = {"PG", "SG", "SF", "PF", "C", "G", "F", "G-F", "F-G", "F-C", "C-F"}

    OUTPUT_FILE = "nba_trade_engine_data.xlsx"



    def make_session():

        session = requests.Session()

        session.headers.update({"User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/139.0.0.0 Safari/537.36", "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,image/avif,image/webp,*/*;q=0.8", "Accept-Language": "en-US,en;q=0.9", "Referer": "https://www.google.com/"})

        return session



    def get_html(session, url, allow_insecure_fallback = False):

        try:

            response = session.get(url, timeout = 30)

        except requests.exceptions.SSLError:

            if not allow_insecure_fallback:

                raise

            print(f"    SSL verification failed for {url}. Retrying this public page with verify = False...")

            response = session.get(url, timeout = 30, verify = False)

        if response.status_code == 403:

            raise RuntimeError(f"403 Forbidden from {url}. The source may be blocking automated requests; rerun later or increase REQUEST_DELAY.")

        response.raise_for_status()

        return response.text



    def flatten_columns(df):

        df = df.copy()

        if isinstance(df.columns, pd.MultiIndex):

            df.columns = [" ".join([str(x) for x in col if str(x) != "nan" and not str(x).startswith("Unnamed")]).strip() for col in df.columns]

        else:

            df.columns = [str(col).strip() for col in df.columns]

        return df



    def read_tables(html):

        try:

            return [flatten_columns(df) for df in pd.read_html(StringIO(html))]

        except ValueError:

            return []



    def money_to_number(value):

        if pd.isna(value):

            return pd.NA

        text = str(value).strip().replace(",", "")

        if text in {"", "-", "—", "nan", "None"}:

            return pd.NA

        match = re.search(r"\$?(-?\d+(?:\.\d+)?)\s*([MBK])?", text, flags = re.I)

        if not match:

            return pd.NA

        number = float(match.group(1))

        suffix = (match.group(2) or "").upper()

        if suffix == "B":

            number *= 1_000_000_000

        elif suffix == "M":

            number *= 1_000_000

        elif suffix == "K":

            number *= 1_000

        return int(round(number))



    def percent_to_number(value):

        if pd.isna(value):

            return pd.NA

        matches = re.findall(r"(-?\d+(?:\.\d+)?)%", str(value))

        return float(matches[-1]) if matches else pd.NA



    def normalize_name(value):

        value = re.sub(r"\s+", " ", str(value)).strip()

        value = re.sub(r"\s+[†‡*]+$", "", value).strip()

        value = re.sub(r"\s*\([^)]{1,8}\)\s*$", "", value).strip()

        value = re.sub(r",\s*[A-Z/-]{1,6}$", "", value, flags = re.I).strip()

        tokens = value.split()

        for prefix_len in range(1, (len(tokens) // 2) + 1):

            if tokens[:prefix_len] == tokens[-prefix_len:] and len(tokens[prefix_len:]) >= 2:

                value = " ".join(tokens[prefix_len:])

                break

        return value



    def season_start_year(label):

        match = re.match(r"^(20\d{2})-\d{2}$", str(label).strip())

        return int(match.group(1)) if match else None



    def find_table(tables, required_columns):

        required_columns = [x.lower() for x in required_columns]

        for df in tables:

            columns = [str(col).lower() for col in df.columns]

            if all(any(required in col for col in columns) for required in required_columns):

                return df.copy()

        return None



    def table_header_cells(table):

        for tr in table.find_all("tr"):

            cells = tr.find_all(["th", "td"])

            texts = [re.sub(r"\s+", " ", cell.get_text(" ", strip = True)).strip() for cell in cells]

            if any(text.lower().startswith("player") for text in texts) and any(season_start_year(text) is not None for text in texts):

                return texts

        return []



    def find_html_table(soup, required_headers):

        required_headers = [x.lower() for x in required_headers]

        for table in soup.find_all("table"):

            headers = [re.sub(r"\s+", " ", th.get_text(" ", strip = True)).strip() for th in table.find_all("th")]

            lower = [header.lower() for header in headers]

            if all(any(required in header for header in lower) for required in required_headers):

                return table

        return None



    def find_table_after_heading(soup, heading_text, required_headers = None):

        heading = soup.find(lambda tag: tag.name in {"h1", "h2", "h3", "h4"} and heading_text.lower() in re.sub(r"\s+", " ", tag.get_text(" ", strip = True)).lower())

        if heading is not None:

            table = heading.find_next("table")

            if table is not None:

                if required_headers is None:

                    return table

                headers = [re.sub(r"\s+", " ", th.get_text(" ", strip = True)).strip().lower() for th in table.find_all("th")]

                if all(any(required.lower() in header for header in headers) for required in required_headers):

                    return table

        return find_html_table(soup, required_headers or [])



    def clean_player_from_cell(cell):

        for anchor in cell.find_all("a"):

            href = str(anchor.get("href", ""))

            text = normalize_name(anchor.get_text(" ", strip = True))

            if text and ("/nba/player" in href or "/redirect/player" in href):

                return text

        return normalize_name(cell.get_text(" ", strip = True))



    def get_spotrac_contract_terms(session):

        rows = []

        for i, (team_slug, (team_name, team_abbr)) in enumerate(SPOTRAC_TEAMS.items(), start = 1):

            print(f"[{i:02d}/30] Spotrac contracts: {team_name}")

            try:

                html = get_html(session, f"https://www.spotrac.com/nba/{team_slug}/contracts")

                soup = BeautifulSoup(html, "html.parser")

                table = find_html_table(soup, ["player", "start", "end", "yrs", "value", "aav"])

                if table is None:

                    print("    Contract table not found.")

                    continue

                headers = []

                header_row = None

                for tr in table.find_all("tr"):

                    texts = [re.sub(r"\s+", " ", cell.get_text(" ", strip = True)).strip() for cell in tr.find_all(["th", "td"])]

                    if any(text.lower().startswith("player") for text in texts) and any(text.lower() == "yrs" for text in texts):

                        headers = texts

                        header_row = tr

                        break

                if not headers:

                    print("    Contract headers not found.")

                    continue

                header_map = {re.sub(r"\s*\(\d+\)\s*$", "", header).strip(): idx for idx, header in enumerate(headers)}

                for tr in header_row.find_all_next("tr"):

                    if tr.find_parent("table") != table:

                        break

                    cells = tr.find_all("td")

                    if len(cells) < max(6, len(headers) - 2):

                        continue

                    name = clean_player_from_cell(cells[0])

                    if not name or name.lower() in {"nan", "totals", "player"}:

                        continue

                    def cell_text(header_name):

                        idx = header_map.get(header_name)

                        return cells[idx].get_text(" ", strip = True) if idx is not None and idx < len(cells) else pd.NA

                    rows.append({"Name": name, "team": team_name, "team_abbreviation": team_abbr, "position": cell_text("Pos"), "signed_year": pd.to_numeric(cell_text("Start Year"), errors = "coerce"), "contract_type": cell_text("Type"), "age_at_signing": pd.to_numeric(cell_text("Age At Signing"), errors = "coerce"), "contract_start": pd.to_numeric(cell_text("Start"), errors = "coerce"), "contract_end": pd.to_numeric(cell_text("End"), errors = "coerce"), "contract_years": pd.to_numeric(cell_text("Yrs"), errors = "coerce"), "contract_value": money_to_number(cell_text("Value")), "contract_aav": money_to_number(cell_text("AAV")), "guaranteed_at_signing": money_to_number(cell_text("GTD @ Sign")), "practical_guaranteed": money_to_number(cell_text("Practical GTD"))})

            except Exception as e:

                print(f"    ERROR: {e}")

            time.sleep(REQUEST_DELAY)

        return pd.DataFrame(rows)



    def parse_deadlines(tables):

        table = find_table(tables, ["Deadline Date", "Player", "Type", "Value"])

        rows = []

        if table is None:

            return pd.DataFrame(rows)

        for _, row in table.iterrows():

            name = normalize_name(row.get("Player"))

            type_text = str(row.get("Type", "")).strip()

            season_match = re.search(r"(20\d{2}-\d{2})", type_text)

            if not name or season_match is None:

                continue

            upper = type_text.upper()

            rows.append({"Name": name, "season": season_match.group(1), "player_option": upper.startswith("PLAYER ") and "PLAYER OPTION" in upper, "team_option": (upper.startswith("CLUB ") or upper.startswith("TEAM ")) and ("CLUB OPTION" in upper or "TEAM OPTION" in upper), "qualifying_offer": "RFA / QO" in upper or "QUALIFYING OFFER" in upper, "extension_eligible": "EXTENSION ELIGIBLE" in upper, "guarantee_deadline": "GUARANTEED" in upper, "deadline_date": pd.to_datetime(row.get("Deadline Date"), errors = "coerce"), "deadline_value": money_to_number(row.get("Value")), "deadline_type": type_text})

        return pd.DataFrame(rows)



    def get_spotrac_yearly_contracts_and_cap(session):

        contract_rows = []

        cap_rows = []

        deadline_frames = []

        for i, (team_slug, (team_name, team_abbr)) in enumerate(SPOTRAC_TEAMS.items(), start = 1):

            print(f"[{i:02d}/30] Spotrac multi-year cap: {team_name}")

            try:

                html = get_html(session, f"https://www.spotrac.com/nba/{team_slug}/yearly/")

                soup = BeautifulSoup(html, "html.parser")

                tables = read_tables(html)

                deadlines = parse_deadlines(tables)

                if not deadlines.empty:

                    deadlines["team"] = team_name

                    deadlines["team_abbreviation"] = team_abbr

                    deadline_frames.append(deadlines)

                active_table = find_table_after_heading(soup, "Active Roster", ["player", "pos", "age", "2026-27"])

                if active_table is not None:

                    headers = table_header_cells(active_table)

                    if headers:

                        header_row = None

                        for tr in active_table.find_all("tr"):

                            texts = [re.sub(r"\s+", " ", cell.get_text(" ", strip = True)).strip() for cell in tr.find_all(["th", "td"])]

                            if texts == headers:

                                header_row = tr

                                break

                        for tr in header_row.find_all_next("tr") if header_row is not None else []:

                            if tr.find_parent("table") != active_table:

                                break

                            cells = tr.find_all("td")

                            if len(cells) < len(headers):

                                continue

                            name = clean_player_from_cell(cells[0])

                            if not name or name.lower() in {"nan", "totals", "player"}:

                                continue

                            position = cells[1].get_text(" ", strip = True) if len(cells) > 1 else pd.NA

                            age = pd.to_numeric(cells[2].get_text(" ", strip = True), errors = "coerce") if len(cells) > 2 else pd.NA

                            for idx, column in enumerate(headers):

                                year = season_start_year(column)

                                if year is None or year < MIN_CONTRACT_SEASON or idx >= len(cells):

                                    continue

                                raw = re.sub(r"\s+", " ", cells[idx].get_text(" ", strip = True)).strip()

                                if raw in {"", "nan", "-", "—"} or "UFA" in raw.upper() or "RFA" in raw.upper():

                                    continue

                                salary = money_to_number(raw)

                                if pd.isna(salary):

                                    continue

                                contract_rows.append({"Name": name, "team": team_name, "team_abbreviation": team_abbr, "position": position, "age": age, "season": column, "season_start": year, "salary": salary, "cap_pct": percent_to_number(raw), "player_option": False, "team_option": False, "qualifying_offer": False, "extension_eligible": "ext. elig" in raw.lower(), "salary_status_raw": raw})

                summary = None

                for table in tables:

                    first_col = table.iloc[:, 0].astype(str).str.strip() if not table.empty else pd.Series(dtype = str)

                    if first_col.str.contains("Cap Maximum", case = False, na = False).any() and first_col.str.contains("1st Apron", case = False, na = False).any():

                        summary = table.copy()

                        break

                if summary is not None:

                    label_col = summary.columns[0]

                    labels = summary[label_col].astype(str).str.strip().tolist()

                    def row_values(label, start_index = 0):

                        for idx in range(start_index, len(labels)):

                            if labels[idx].lower() == label.lower():

                                return idx, summary.iloc[idx]

                        return None, None

                    _, cap_max_row = row_values("Cap Maximum")

                    _, active_row = row_values("Active Cap")

                    _, dead_row = row_values("Dead Cap")

                    _, holds_row = row_values("Cap Holds")

                    _, total_row = row_values("Total Cap Allocations")

                    _, space_row = row_values("Cap Space")

                    first_header = next((idx for idx, label in enumerate(labels) if "1st Apron" in label and "Space" not in label), None)

                    second_header = next((idx for idx, label in enumerate(labels) if "2nd Apron" in label and "Space" not in label), None)

                    _, first_threshold_row = row_values("Threshold", first_header + 1 if first_header is not None else 0)

                    _, first_alloc_row = row_values("Allocations", first_header + 1 if first_header is not None else 0)

                    _, first_space_row = row_values("1st Apron Space", first_header + 1 if first_header is not None else 0)

                    _, second_threshold_row = row_values("Threshold", second_header + 1 if second_header is not None else 0)

                    _, second_alloc_row = row_values("Allocations", second_header + 1 if second_header is not None else 0)

                    _, second_space_row = row_values("2nd Apron Space", second_header + 1 if second_header is not None else 0)

                    for column in summary.columns[1:]:

                        year = season_start_year(column)

                        if year is None or year < MIN_CONTRACT_SEASON:

                            continue

                        cap_space = money_to_number(space_row.get(column)) if space_row is not None else pd.NA

                        first_space = money_to_number(first_space_row.get(column)) if first_space_row is not None else pd.NA

                        second_space = money_to_number(second_space_row.get(column)) if second_space_row is not None else pd.NA

                        cap_rows.append({"team": team_name, "team_abbreviation": team_abbr, "season": column, "season_start": year, "salary_cap": money_to_number(cap_max_row.get(column)) if cap_max_row is not None else pd.NA, "active_cap": money_to_number(active_row.get(column)) if active_row is not None else pd.NA, "dead_cap": money_to_number(dead_row.get(column)) if dead_row is not None else pd.NA, "cap_holds": money_to_number(holds_row.get(column)) if holds_row is not None else pd.NA, "total_cap_allocations": money_to_number(total_row.get(column)) if total_row is not None else pd.NA, "cap_space": cap_space, "first_apron": money_to_number(first_threshold_row.get(column)) if first_threshold_row is not None else pd.NA, "first_apron_allocations": money_to_number(first_alloc_row.get(column)) if first_alloc_row is not None else pd.NA, "first_apron_space": first_space, "second_apron": money_to_number(second_threshold_row.get(column)) if second_threshold_row is not None else pd.NA, "second_apron_allocations": money_to_number(second_alloc_row.get(column)) if second_alloc_row is not None else pd.NA, "second_apron_space": second_space, "above_cap": bool(cap_space < 0) if not pd.isna(cap_space) else pd.NA, "above_first_apron": bool(first_space < 0) if not pd.isna(first_space) else pd.NA, "above_second_apron": bool(second_space < 0) if not pd.isna(second_space) else pd.NA})

            except Exception as e:

                print(f"    ERROR: {e}")

            time.sleep(REQUEST_DELAY)

        contracts_df = pd.DataFrame(contract_rows)

        deadlines_df = pd.concat(deadline_frames, ignore_index = True) if deadline_frames else pd.DataFrame()

        if not contracts_df.empty and not deadlines_df.empty:

            flag_cols = ["player_option", "team_option", "qualifying_offer", "extension_eligible", "guarantee_deadline"]

            grouped = deadlines_df.groupby(["Name", "team", "team_abbreviation", "season"], as_index = False)[flag_cols].max()

            contracts_df = contracts_df.merge(grouped, on = ["Name", "team", "team_abbreviation", "season"], how = "left", suffixes = ("", "_deadline"))

            for col in ["player_option", "team_option", "qualifying_offer", "extension_eligible"]:

                deadline_col = f"{col}_deadline"

                if deadline_col in contracts_df.columns:

                    contracts_df[col] = contracts_df[col].fillna(False) | contracts_df[deadline_col].fillna(False)

                    contracts_df = contracts_df.drop(columns = deadline_col)

            if "guarantee_deadline" in contracts_df.columns:

                contracts_df["guarantee_deadline"] = contracts_df["guarantee_deadline"].fillna(False).astype(bool)

        return contracts_df, pd.DataFrame(cap_rows), deadlines_df



    def get_spotrac_trade_exceptions(session):

        rows = []

        for i, (team_slug, (team_name, team_abbr)) in enumerate(SPOTRAC_TEAMS.items(), start = 1):

            print(f"[{i:02d}/30] Spotrac trade exceptions: {team_name}")

            try:

                html = get_html(session, f"https://www.spotrac.com/nba/{team_slug}/cap/_/year/{MIN_CONTRACT_SEASON}")

                table = find_table(read_tables(html), ["Type", "Reason", "Expires", "Original", "Available"])

                if table is None:

                    continue

                for _, row in table.iterrows():

                    exception_type = str(row.get("Type", "")).strip()

                    if exception_type.lower() != "trade":

                        continue

                    rows.append({"team": team_name, "team_abbreviation": team_abbr, "season": f"{MIN_CONTRACT_SEASON}-{str(MIN_CONTRACT_SEASON + 1)[-2:]}", "exception_type": exception_type, "reason": row.get("Reason"), "used_on": row.get("Used On"), "expires": pd.to_datetime(row.get("Expires"), errors = "coerce"), "original_amount": money_to_number(row.get("Original")), "available_amount": money_to_number(row.get("Available"))})

            except Exception as e:

                print(f"    ERROR: {e}")

            time.sleep(REQUEST_DELAY)

        return pd.DataFrame(rows)



    def classify_transaction(description):

        text = str(description).lower()

        if "extension" in text or "extended" in text:

            return "Extension"

        if "traded" in text or "trade to" in text or "trade from" in text:

            return "Trade"

        if "signed" in text or "re-signed" in text:

            return "Signing"

        return None



    def parse_spotrac_transaction_page(html, source_team, source_abbr):

        soup = BeautifulSoup(html, "html.parser")

        date_pattern = re.compile(r"\b(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)\s+\d{1,2},\s+20\d{2}\b")

        keyword_pattern = re.compile(r"\b(?:signed|re-signed|extension|extended|traded)\b", flags = re.I)

        containers = []

        for tag in soup.find_all("li"):

            text = " ".join(tag.stripped_strings)

            if date_pattern.search(text) and keyword_pattern.search(text):

                containers.append(tag)

        if not containers:

            candidates = []

            for tag in soup.find_all("div"):

                text = " ".join(tag.stripped_strings)

                if date_pattern.search(text) and keyword_pattern.search(text) and 20 <= len(text) <= 1200:

                    child_match = any(date_pattern.search(" ".join(child.stripped_strings)) and keyword_pattern.search(" ".join(child.stripped_strings)) for child in tag.find_all("div", recursive = False))

                    if not child_match:

                        candidates.append(tag)

            containers = candidates

        rows = []

        seen = set()

        for container in containers:

            text = re.sub(r"\s+", " ", " ".join(container.stripped_strings)).strip()

            date_match = date_pattern.search(text)

            if not date_match:

                continue

            player_links = [a for a in container.find_all("a") if "/nba/player" in str(a.get("href", "")) or "/redirect/player" in str(a.get("href", ""))]

            if player_links:

                name = normalize_name(player_links[0].get_text(" ", strip = True))

            else:

                after_date = text[date_match.end():].strip()

                name_match = re.match(r"([^,]+?)(?:,\s*[A-Z]{1,3}|\s*\([A-Z/-]{1,6}\))", after_date)

                name = normalize_name(name_match.group(1)) if name_match else normalize_name(text[:date_match.start()])

            if not name or name.lower() == "nan":

                continue

            position_match = re.search(re.escape(name) + r"\s*(?:,\s*|\()([A-Z/-]{1,6})\)?", text, flags = re.I)

            position = position_match.group(1).upper() if position_match else pd.NA

            if not pd.isna(position) and position not in PLAYER_POSITIONS:

                continue

            description = text[date_match.end():].strip().lstrip("-").strip()

            description = re.sub(r"^" + re.escape(name) + r"\s*(?:,\s*[A-Z/-]{1,6}|\s*\([A-Z/-]{1,6}\))?\s*", "", description, flags = re.I)

            description = re.sub(r"^(Signing|Trade|Extension)\s+\1\s+", r"\1 ", description, flags = re.I)

            if "pending" in description.lower():

                continue

            transaction_type = classify_transaction(description)

            if transaction_type is None:

                continue

            transaction_date = pd.to_datetime(date_match.group(0), errors = "coerce")

            key = (name, transaction_date, transaction_type, description)

            if key in seen:

                continue

            seen.add(key)

            rows.append({"Name": name, "position": position, "date": transaction_date, "transaction_type": transaction_type, "team": source_team, "team_abbreviation": source_abbr, "description": description})

        return rows



    def get_spotrac_transactions(session):

        rows = []

        end_date = date.today().isoformat()

        for i, (_, (team_name, team_abbr)) in enumerate(SPOTRAC_TEAMS.items(), start = 1):

            print(f"[{i:02d}/30] Spotrac transactions: {team_name}")

            try:

                url = f"https://www.spotrac.com/nba/transactions/_/start/{TRANSACTION_START}/end/{end_date}/team/{team_abbr.lower()}"

                rows.extend(parse_spotrac_transaction_page(get_html(session, url), team_name, team_abbr))

            except Exception as e:

                print(f"    ERROR: {e}")

            time.sleep(REQUEST_DELAY)

        df = pd.DataFrame(rows)

        if not df.empty:

            df = df.drop_duplicates(subset = ["Name", "date", "transaction_type", "description"]).sort_values(["date", "Name"], ascending = [False, True]).reset_index(drop = True)

        return df



    def pick_is_referenced(traded_df, team_name, draft_year, draft_round):

        round_terms = "(?:1st|first)" if draft_round == 1 else "(?:2nd|second)"

        combined = (traded_df["asset"].fillna("") + " " + traded_df["details"].fillna("")).astype(str)

        for alias in PICK_ALIASES.get(team_name, [team_name]):

            possessive = re.escape(alias) + r"(?:'s|')"

            pattern = re.compile(rf"{possessive}\s+{draft_year}\s+{round_terms}\s+round\s+pick", flags = re.I)

            if combined.str.contains(pattern, regex = True, na = False).any():

                return True

        return False



    def get_courtside_draft_assets(session):

        print("Pulling Courtside News future draft assets...")

        url = "https://courtsidenews.com/nba/draft/picks"

        html = get_html(session, url, allow_insecure_fallback = True)

        soup = BeautifulSoup(html, "html.parser")

        strings = [re.sub(r"\s+", " ", text).strip() for text in soup.stripped_strings]

        team_lookup = {team_name: (team_name, team_abbr) for _, (team_name, team_abbr) in SPOTRAC_TEAMS.items()}

        team_lookup["LA Clippers"] = ("Los Angeles Clippers", "LAC")

        team_names = set(team_lookup)

        rows = []

        current_team = None

        current_team_abbr = None

        current_year = None

        current_direction = None

        current_row = None

        started = False

        def finish_current():

            nonlocal current_row

            if current_row is None:

                return

            details = " ".join(current_row.pop("detail_parts", [])).strip()

            current_row["details"] = details if details else pd.NA

            combined = f"{current_row.get('asset', '')} {details}".lower()

            current_row["is_swap"] = "swap" in combined or "right to swap" in combined

            current_row["is_protected"] = "protect" in combined

            current_row["is_conditional"] = any(term in combined for term in ["if ", "more favorable", "less favorable", "most favorable", "least favorable", "convey", "extinguished", "condition", "outgoing"])

            rows.append(current_row)

            current_row = None

        for text in strings:

            if text == "NBA Future Draft Picks Tracker":

                started = True

                continue

            if not started:

                continue

            if text == "NBA Standings":

                finish_current()

                break

            if text in team_names:

                finish_current()

                current_team, current_team_abbr = team_lookup[text]

                current_year = None

                current_direction = None

                continue

            if current_team is None:

                continue

            if re.fullmatch(r"20\d{2}", text):

                year = int(text)

                if year >= MIN_CONTRACT_SEASON + 1:

                    finish_current()

                    current_year = year

                    current_direction = None

                continue

            if text in {"IN", "OUT"}:

                finish_current()

                current_direction = text

                continue

            asset_match = re.match(r"^(20\d{2}) (first|second) round draft pick\b", text, flags = re.I)

            if asset_match and current_year is not None and current_direction is not None:

                finish_current()

                asset = re.sub(r"\s*Protections and trade notes\s*$", "", text, flags = re.I).strip()

                current_row = {"team": current_team, "team_abbreviation": current_team_abbr, "draft_year": int(asset_match.group(1)), "round": 1 if asset_match.group(2).lower() == "first" else 2, "direction": current_direction, "asset": asset, "detail_parts": []}

                continue

            if current_row is not None and text.lower() != "protections and trade notes":

                current_row["detail_parts"].append(text)

        finish_current()

        traded_df = pd.DataFrame(rows)

        if traded_df.empty:

            raise RuntimeError("No future draft-pick data was parsed from Courtside News.")

        traded_df = traded_df[traded_df["draft_year"] >= MIN_CONTRACT_SEASON + 1].copy()

        own_rows = []

        max_draft_year = max(int(traded_df["draft_year"].max()), MIN_CONTRACT_SEASON + 7)

        for _, (team_name, team_abbr) in SPOTRAC_TEAMS.items():

            for draft_year in range(MIN_CONTRACT_SEASON + 1, max_draft_year + 1):

                for draft_round in [1, 2]:

                    if not pick_is_referenced(traded_df, team_name, draft_year, draft_round):

                        round_name = "first" if draft_round == 1 else "second"

                        own_rows.append({"team": team_name, "team_abbreviation": team_abbr, "draft_year": draft_year, "round": draft_round, "direction": "OWN", "asset": f"{draft_year} {round_name} round own draft pick", "details": "Unencumbered own pick; this original pick is not referenced in any traded-pick obligation, protection, or swap in the tracker.", "is_swap": False, "is_protected": False, "is_conditional": False})

        own_df = pd.DataFrame(own_rows)

        draft_assets_df = pd.concat([traded_df, own_df], ignore_index = True, sort = False)

        direction_order = pd.CategoricalDtype(["OWN", "IN", "OUT"], ordered = True)

        draft_assets_df["direction"] = draft_assets_df["direction"].astype(direction_order)

        draft_assets_df = draft_assets_df.sort_values(["team", "draft_year", "round", "direction", "asset"]).reset_index(drop = True)

        draft_assets_df["direction"] = draft_assets_df["direction"].astype(str)

        return draft_assets_df



    def attach_contract_terms(annual_contracts_df, contract_terms_df):

        if annual_contracts_df.empty or contract_terms_df.empty:

            return annual_contracts_df.copy()

        annual = annual_contracts_df.copy().reset_index(drop = True)

        terms = contract_terms_df.copy()

        term_columns = ["signed_year", "contract_type", "age_at_signing", "contract_start", "contract_end", "contract_years", "contract_value", "contract_aav", "guaranteed_at_signing", "practical_guaranteed"]

        output_rows = []

        for _, row in annual.iterrows():

            candidates = terms[(terms["Name"] == row["Name"]) & (terms["team_abbreviation"] == row["team_abbreviation"])].copy()

            season_year = int(row["season_start"])

            if not candidates.empty:

                valid = candidates[(pd.to_numeric(candidates["contract_start"], errors = "coerce") <= season_year) & (pd.to_numeric(candidates["contract_end"], errors = "coerce") >= season_year)]

                if valid.empty:

                    valid = candidates[pd.to_numeric(candidates["contract_start"], errors = "coerce") <= season_year]

                if valid.empty:

                    valid = candidates

                valid = valid.assign(_start = pd.to_numeric(valid["contract_start"], errors = "coerce")).sort_values("_start", ascending = False)

                chosen = valid.iloc[0]

                for col in term_columns:

                    row[col] = chosen.get(col, pd.NA)

            else:

                for col in term_columns:

                    row[col] = pd.NA

            output_rows.append(row.to_dict())

        return pd.DataFrame(output_rows)



    def add_transaction_dates_to_contracts(contracts_df, transactions_df):

        if contracts_df.empty or transactions_df.empty:

            return contracts_df

        transactions = transactions_df.copy()

        transactions["Name"] = transactions["Name"].map(normalize_name)

        transactions["event_year"] = pd.to_datetime(transactions["date"], errors = "coerce").dt.year

        contracts = contracts_df.copy()

        signed_dates = []

        signed_types = []

        signed_descriptions = []

        latest_trade_dates = []

        latest_trade_descriptions = []

        for _, contract in contracts.iterrows():

            name = normalize_name(contract["Name"])

            signed_year = pd.to_numeric(contract.get("signed_year"), errors = "coerce")

            events = transactions[(transactions["Name"] == name) & (transactions["transaction_type"].isin(["Signing", "Extension"]))].copy()

            if not pd.isna(signed_year):

                same_year = events[events["event_year"] == int(signed_year)]

                if not same_year.empty:

                    events = same_year

            if not events.empty:

                event = events.sort_values("date").iloc[-1]

                signed_dates.append(event["date"])

                signed_types.append(event["transaction_type"])

                signed_descriptions.append(event["description"])

            else:

                signed_dates.append(pd.NaT)

                signed_types.append(pd.NA)

                signed_descriptions.append(pd.NA)

            trades = transactions[(transactions["Name"] == name) & (transactions["transaction_type"] == "Trade")].copy()

            if not trades.empty:

                trade = trades.sort_values("date").iloc[-1]

                latest_trade_dates.append(trade["date"])

                latest_trade_descriptions.append(trade["description"])

            else:

                latest_trade_dates.append(pd.NaT)

                latest_trade_descriptions.append(pd.NA)

        contracts["contract_signed_date"] = signed_dates

        contracts["contract_event_type"] = signed_types

        contracts["contract_event_description"] = signed_descriptions

        contracts["latest_trade_date"] = latest_trade_dates

        contracts["latest_trade_description"] = latest_trade_descriptions

        return contracts



    def validate_data(contracts_df, team_cap_df, draft_assets_df):

        errors = []

        if contracts_df.empty:

            errors.append("contracts_df is empty")

        else:

            bad_salary = contracts_df[pd.to_numeric(contracts_df["salary"], errors = "coerce") > 120_000_000]

            if not bad_salary.empty:

                errors.append(f"{len(bad_salary)} contract rows have salary > $120M; salary/percentage parsing likely failed")

            bad_pct = contracts_df[pd.to_numeric(contracts_df["cap_pct"], errors = "coerce") > 60]

            if not bad_pct.empty:

                errors.append(f"{len(bad_pct)} contract rows have cap_pct > 60%; percentage parsing likely failed")

            current = contracts_df[contracts_df["season_start"] == MIN_CONTRACT_SEASON]

            counts = current.groupby("team_abbreviation")["Name"].nunique()

            missing_or_tiny = [abbr for _, (_, abbr) in SPOTRAC_TEAMS.items() if counts.get(abbr, 0) < 5]

            if missing_or_tiny:

                errors.append(f"Current-season contract pull returned fewer than 5 players for: {', '.join(missing_or_tiny)}")

        cap_current = team_cap_df[team_cap_df["season_start"] == MIN_CONTRACT_SEASON] if not team_cap_df.empty else pd.DataFrame()

        if cap_current.empty or cap_current["team_abbreviation"].nunique() != 30:

            errors.append("Current-season cap table does not contain all 30 teams")

        if draft_assets_df.empty:

            errors.append("draft_assets_df is empty")

        if errors:

            raise RuntimeError("DATA VALIDATION FAILED:\n- " + "\n- ".join(errors))

        print("Data validation passed.")



    def get_nba_trade_engine_data():

        session = make_session()

        draft_assets_df = get_courtside_draft_assets(session)

        print(f"Draft assets successfully pulled: {len(draft_assets_df):,} rows\n")

        contract_terms_df = get_spotrac_contract_terms(session)

        annual_contracts_df, team_cap_df, deadlines_df = get_spotrac_yearly_contracts_and_cap(session)

        trade_exceptions_df = get_spotrac_trade_exceptions(session)

        transactions_df = get_spotrac_transactions(session)

        contracts_df = attach_contract_terms(annual_contracts_df, contract_terms_df) if not annual_contracts_df.empty else contract_terms_df.copy()

        contracts_df = add_transaction_dates_to_contracts(contracts_df, transactions_df)

        if not contracts_df.empty and "season" in contracts_df.columns:

            contracts_df = contracts_df.drop_duplicates(subset = ["Name", "team", "season"]).sort_values(["team", "Name", "season"]).reset_index(drop = True)

        validate_data(contracts_df, team_cap_df, draft_assets_df)

        return contracts_df, transactions_df, team_cap_df, trade_exceptions_df, draft_assets_df, deadlines_df




    def create_team_players_dictionary(contracts_df):
        team_players = {
            team_name: []
            for _, (team_name, _) in SPOTRAC_TEAMS.items()
        }

        if contracts_df.empty:
            return team_players

        current_contracts = contracts_df[
            contracts_df["season_start"] == MIN_CONTRACT_SEASON
        ].copy()

        for team_name in team_players:
            players = (
                current_contracts.loc[
                    current_contracts["team"] == team_name,
                    "Name"
                ]
                .dropna()
                .drop_duplicates()
                .sort_values()
                .tolist()
            )

            team_players[team_name] = players

        return team_players


    def save_team_workbook(contracts_df, transactions_df, team_cap_df, trade_exceptions_df, draft_assets_df, deadlines_df, output_file = OUTPUT_FILE):

        sections = [("CONTRACTS", contracts_df), ("TRANSACTIONS", transactions_df), ("TEAM CAP / APRONS", team_cap_df), ("TRADE EXCEPTIONS", trade_exceptions_df), ("CONTRACT DEADLINES", deadlines_df), ("DRAFT ASSETS", draft_assets_df)]

        with pd.ExcelWriter(output_file, engine = "openpyxl", datetime_format = "yyyy-mm-dd") as writer:

            for _, (team_name, team_abbr) in SPOTRAC_TEAMS.items():

                sheet_name = team_name[:31]

                start_row = 0

                section_title_rows = []

                for section_name, df in sections:

                    section_title_rows.append(start_row + 1)

                    pd.DataFrame([[section_name]]).to_excel(writer, sheet_name = sheet_name, startrow = start_row, index = False, header = False)

                    start_row += 1

                    if df.empty:

                        team_df = pd.DataFrame()

                    elif "team_abbreviation" in df.columns:

                        team_df = df[df["team_abbreviation"] == team_abbr].copy()

                    elif "team" in df.columns:

                        team_df = df[df["team"] == team_name].copy()

                    else:

                        team_df = pd.DataFrame()

                    if team_df.empty:

                        pd.DataFrame({"Info": ["No data found"]}).to_excel(writer, sheet_name = sheet_name, startrow = start_row, index = False)

                        start_row += 3

                    else:

                        team_df.to_excel(writer, sheet_name = sheet_name, startrow = start_row, index = False)

                        start_row += len(team_df) + 3

                worksheet = writer.book[sheet_name]

                worksheet.freeze_panes = "A2"

                for row_number in section_title_rows:

                    cell = worksheet.cell(row = row_number, column = 1)

                    cell.font = cell.font.copy(bold = True, size = 12)

                for row in worksheet.iter_rows():

                    for cell in row:

                        if isinstance(cell.value, pd.Timestamp):

                            cell.number_format = "yyyy-mm-dd"

                for column_cells in worksheet.columns:

                    max_length = 0

                    column_letter = column_cells[0].column_letter

                    for cell in column_cells:

                        if cell.value is not None:

                            max_length = max(max_length, len(str(cell.value)))

                    worksheet.column_dimensions[column_letter].width = min(max(max_length + 2, 10), 40)

        print(f"Saved Excel workbook: {output_file}")

    # ============================================================
    # STAGE 1 — EXECUTE DATA COLLECTION AND SAVE TEAM WORKBOOK
    # ============================================================
    urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)
    contracts_df, transactions_df, team_cap_df, trade_exceptions_df, draft_assets_df, deadlines_df = get_nba_trade_engine_data()
    team_players = create_team_players_dictionary(contracts_df)
    print("\nCOLLECTION COMPLETE")
    print(f"Contracts: {len(contracts_df):,}; Transactions: {len(transactions_df):,}; Teams: {len(team_players)}")
    if SAVE_EXCEL:
        save_team_workbook(contracts_df, transactions_df, team_cap_df, trade_exceptions_df, draft_assets_df, deadlines_df)
    else:
        # Stage 2 reads the workbook, so it must exist for the current run.
        raise ValueError("SAVE_EXCEL must be True for this workbook-based pipeline")

    # ============================================================
    # STAGE 2 — LOAD PLAYER NAMES AND EVALUATE PLAYERS USING GPT
    # ============================================================
    def get_league_players() -> dict[str, list[str]]:
        FILE_NAME = OUTPUT_FILE
        CURRENT_SEASON = "2026-27"

        workbook = load_workbook(
            FILE_NAME,
            read_only=True,
            data_only=True
        )

        league_players = {}

        try:
            for sheet in workbook.worksheets:
                players = set()
                in_contracts = False
                headers = None

                for row in sheet.iter_rows(values_only=True):
                    first_cell = row[0]

                    # Locate the contracts section.
                    if first_cell == "CONTRACTS":
                        in_contracts = True
                        continue

                    if not in_contracts:
                        continue

                    # Stop at the end of the contracts section.
                    if first_cell is None:
                        break

                    # Read column positions from the header.
                    if headers is None:
                        headers = {
                            str(value).strip(): index
                            for index, value in enumerate(row)
                            if value is not None
                        }
                        continue

                    name = row[headers["Name"]]
                    season = row[headers["season"]]

                    # Only include current-season players.
                    if (
                        isinstance(name, str)
                        and name.strip()
                        and season == CURRENT_SEASON
                    ):
                        players.add(name.strip())

                league_players[sheet.title] = sorted(players)

        finally:
            workbook.close()

        return league_players
    league_players = get_league_players()
    client = OpenAI(api_key=api_key, http_client=cast(Any, httpx.Client()))
    BATCH_SIZE = 5

    teams = list(league_players.items())

    all_results = []

    total_batches = (len(teams) + BATCH_SIZE - 1) // BATCH_SIZE


    for start_index in range(0, len(teams), BATCH_SIZE):

        batch_items = teams[start_index:start_index + BATCH_SIZE]

        batch = dict(batch_items)

        batch_text = json.dumps(
            batch,
            ensure_ascii=False,
            indent=2
        )

        batch_number = (start_index // BATCH_SIZE) + 1

        print(f"Player evaluation batch {batch_number}/{total_batches}")

        response = client.responses.create(
            model="gpt-5.6",
            tools=[
                {
                    "type": "web_search"
                }
            ],
            input=f"""
    Below is a list of current NBA players grouped by team.

    {batch_text}

    Evaluate EVERY player listed.

    For each player, return:

    1. Attributes:
       - Up to 8 concise attributes describing the player physically,
         mentally, stylistically, or by career stage.
       - Examples include:
         age,
         young,
         veteran,
         speed,
         quickness,
         strength,
         size,
         vertical athleticism,
         durability,
         temperament,
         basketball IQ,
         playoff experience,
         playoff riser,
         clutch,
         inconsistent,
         choker,
         high-motor,
         injury-prone,
         etc.

    2. Strengths:
       - Up to 5 meaningful CURRENT basketball strengths.
       - Possible areas include:
         shooting,
         three-point shooting,
         midrange shooting,
         finishing,
         rim finishing,
         passing,
         playmaking,
         ball handling,
         shot creation,
         off-ball movement,
         screening,
         rebounding,
         offensive rebounding,
         perimeter defense,
         interior defense,
         rim protection,
         switchability,
         athleticism,
         basketball IQ,
         decision making,
         transition offense,
         transition defense,
         post offense,
         post defense,
         etc.

    3. Weaknesses:
       - Up to 5 meaningful CURRENT basketball weaknesses.
       - Use the same types of basketball categories as the strengths.

    IMPORTANT RULES:

    - Be stringent.
    - Evaluate CURRENT ability only.
    - Do NOT evaluate players based on what they were at their historical peak.
    - Do NOT evaluate them based on what they might become in the future.
    - Do NOT manufacture strengths or weaknesses simply to reach 5.
    - If a player only has 2 clear weaknesses, return 2.
    - If a player only has 3 clear strengths, return 3.
    - Use current information when necessary.
    - Distinguish between reputation and current performance.
    - Do not call someone clutch or a choker unless there is meaningful evidence.
    - Do not infer personality traits without reasonable public evidence.
    - Basketball skills should be concrete enough that another program can
      later use them to evaluate trade preferences.
    - Evaluate every player in the supplied batch.
    - Do not omit players.

    Return the information EXACTLY in this general format:

    Team Name:

    Player Name: Attributes ["attribute 1", "attribute 2", ...], Strengths ["strength 1", "strength 2", ...], Weaknesses ["weakness 1", "weakness 2", ...]
    Player Name 2: Attributes ["attribute 1", "attribute 2", ...], Strengths ["strength 1", "strength 2", ...], Weaknesses ["weakness 1", "weakness 2", ...]

    Next Team Name:

    Player Name: Attributes [...], Strengths [...], Weaknesses [...]
    Player Name 2: Attributes [...], Strengths [...], Weaknesses [...]

    Return ONLY the requested team/player information.

    Do not include:
    - an introduction
    - an explanation
    - methodology
    - citations
    - source lists
    - conclusions
    - disclaimers
    """
        )

        batch_result = response.output_text.strip()

        all_results.append(batch_result)

    # ============================================================
    # STAGE 3 — PARSE THE USER'S TRADE PREFERENCES USING GPT
    # ============================================================
    demands = client.responses.create(
        model="gpt-5.6",
        input=f"""
You are interpreting an NBA trade request for a later player-pool filter.
Read the user statement and extract the following, without inventing names:

Team Name: [user's NBA team]
Required Acquisition Players: [specifically named players the user insists on acquiring]
Preferred Acquisition Players: [named players the user is interested in but does not require]
Acquisition Player Types: [roles, skills, traits of other acceptable acquisitions]
Preferred Outgoing Players: [named players the user explicitly wants to trade away or offers to trade]
Other Allowed Outgoing Players: [named players the user explicitly permits trading, without preferring to move]
Untouchable Players: [named players never to trade]
Untouchable Player Types: [traits of players never to trade]
Players to Avoid Acquiring: [specific names]
Player Types to Avoid Acquiring: [traits, weaknesses, roles to exclude]
Teams to Avoid: [teams the user refuses to deal with]

Distinguish REQUIRED from merely PREFERRED acquisitions. Explicitly naming
someone does not necessarily make acquisition mandatory: follow the wording.
Likewise, "willing to part with" means allowed, not necessarily preferred,
unless the user expresses a desire to move that player.
If a category is absent, put [] rather than speculating.

USER STATEMENT:
{user_statement}

Return only the labeled fields above. No explanations or Markdown.
"""
    )

    print("\nPARSED USER PREFERENCES:\n", demands.output_text)

    # ============================================================
    # STAGE 4 — GPT PLAYER-POOL REDUCTION AND PRIORITIZATION
    # Separate desired/required acquisitions from permitted/preferred exits.
    # No CBA evaluation is performed in this GPT stage.
    # ============================================================
    response = client.responses.create(
        model="gpt-5.6",
        input=f"""
You build four DIFFERENT player pools for an NBA trade search.

USER PREFERENCES:
{demands.output_text}

PLAYER PROFILES AND THEIR CURRENT TEAMS:
{all_results}

Using ONLY the supplied team rosters and player attributes/strengths/weaknesses,
produce EXACTLY FOUR Python dictionaries in the order specified below.
Each dictionary maps exact existing NBA team names to lists of exact existing
player names. No duplicates, fabricated names, or empty teams.

DICTIONARY 1: required_targets
- On OTHER teams, players the user MUST acquire, per explicit mandatory wording.
- Do NOT classify all named players as required automatically: distinguish an
  actual condition from a mere preference or interest.
- Honor excluded teams even if the user names a player on such a team, unless
  the user explicitly grants an exception for that team.
- A specifically required acquisition overrides general skill/weakness dislikes,
  but not an explicit team ban.

DICTIONARY 2: tradeable_targets
- On OTHER teams, players the user is interested in acquiring based on requested
  player types, skills, roles or explicit nonmandatory player interests.
- Only include players who reasonably match, using supplied scouting profiles.
- Respect excluded teams, explicitly unwanted players and unwanted player types.
- Do NOT duplicate players already in required_targets.
- Do not include everyone generally allowable: these must be actual targets.

DICTIONARY 3: wanted_to_trade_away
- From the USER'S OWN team ONLY, players the user expresses a genuine desire
  to trade away (not merely willingness to part with).
- Never include untouchables unless the user explicitly overrides that status.

DICTIONARY 4: can_trade_away
- From the USER'S OWN team ONLY, all other players the user is willing to trade.
- Exclude explicitly untouchable names and players matching protected types.
- Preserve explicit user permission to move a player even if they otherwise
  match a broad protected type (e.g. willing to part with Mikal Bridges but
  protecting other good shooters).
- Include named players the user is willing to part with, even if not eager.
- Include all other nonprotected roster players unless the request clearly
  says ONLY specific players may be traded.
- Do NOT duplicate any player already in wanted_to_trade_away.

These four dictionaries must be pairwise DISJOINT. They describe USER
preferences only, NOT legal trade eligibility. Do not apply salary/CBA rules.
Do NOT use external knowledge or web searches. Do NOT change a player's team.

OUTPUT:
Return exactly four raw Python dict expressions (not a list or variable
assignments), in this strict order:
1. required_targets
2. tradeable_targets
3. wanted_to_trade_away
4. can_trade_away

Example output structure (illustrative names only):
{{"Opposing Team": ["Required Player"]}}
{{"Other Team": ["Interesting Player"]}}
{{"My Team": ["Player To Move"]}}
{{"My Team": ["Other Allowed Player"]}}

No headings, explanation, code fences, or text between dicts.
"""
    )

    def parse_four_dictionaries(text):
        try:
            tree = ast.parse(text.strip(), mode="exec")
        except SyntaxError as exc:
            raise ValueError(f"GPT returned invalid pool syntax:\n{text}") from exc

        values = []
        for node in tree.body:
            if not isinstance(node, ast.Expr):
                raise ValueError(f"Unexpected non-dictionary expression:\n{text}")
            value = ast.literal_eval(node.value)
            if not isinstance(value, dict):
                raise ValueError("Expected exactly four dictionaries")
            values.append(value)
        if len(values) != 4:
            raise ValueError(f"Expected four dictionaries, got {len(values)}:\n{text}")
        return values

    required_targets, tradeable_targets, wanted_to_trade_away, can_trade_away = (
        parse_four_dictionaries(response.output_text)
    )

    # Validate GPT's names against the scraped roster: models can misspell or
    # assign players to the wrong team. Keep only verified membership.
    roster_lookup = {
        team: {normalize_name(name).casefold(): name for name in names}
        for team, names in league_players.items()
    }

    def verify_roster_pool(pool, pool_name):
        clean = {}
        for team, names in pool.items():
            if team not in roster_lookup:
                raise ValueError(f"{pool_name}: unknown team {team!r}")
            if not isinstance(names, list):
                raise ValueError(f"{pool_name}: player values must be lists")
            verified = []
            for name in names:
                actual = roster_lookup[team].get(normalize_name(name).casefold())
                if actual is None:
                    raise ValueError(f"{pool_name}: {name!r} not found on {team}")
                verified.append(actual)
            if verified:
                clean[team] = sorted(set(verified))
        return clean

    required_targets = verify_roster_pool(required_targets, "required_targets")
    tradeable_targets = verify_roster_pool(tradeable_targets, "tradeable_targets")
    wanted_to_trade_away = verify_roster_pool(wanted_to_trade_away, "wanted_to_trade_away")
    can_trade_away = verify_roster_pool(can_trade_away, "can_trade_away")

    # Favor the higher-priority group if GPT accidentally duplicates a player.
    seen = set()
    def deduplicate_pool(pool):
        cleaned = {}
        for team, players in pool.items():
            unique = [name for name in players if (team, name) not in seen]
            seen.update((team, name) for name in unique)
            if unique:
                cleaned[team] = unique
        return cleaned

    required_targets = deduplicate_pool(required_targets)
    tradeable_targets = deduplicate_pool(tradeable_targets)
    wanted_to_trade_away = deduplicate_pool(wanted_to_trade_away)
    can_trade_away = deduplicate_pool(can_trade_away)

    print("\nREQUIRED TARGETS (BEFORE CBA):", required_targets)
    print("\nTRADEABLE TARGETS (BEFORE CBA):", tradeable_targets)
    print("\nWANTED TO TRADE AWAY (BEFORE CBA):", wanted_to_trade_away)
    print("\nCAN TRADE AWAY (BEFORE CBA):", can_trade_away)

    # ============================================================
    # STAGE 5 — INDIVIDUAL CBA ELIGIBILITY PRE-FILTERING
    # Only checks whether an INDIVIDUAL player can be traded today.
    # No salaries, salary aggregation, cap room, or matching mechanisms.
    # ============================================================
    TRADE_DATE = date.today()
    TRADE_DEADLINE = date(2027, 2, 11)  # Official 2026-27 NBA deadline.

    def parse_date(value):
        if value is None or pd.isna(value):
            return None
        parsed = pd.to_datetime(value, errors="coerce")
        return parsed.date() if not pd.isna(parsed) else None

    def eligibility_record(row):
        """Return (eligible, reasons, unknowns), with no salary checks."""
        name = str(row.get("Name", ""))
        team = str(row.get("team", ""))
        reasons, unknowns = [], []
        if TRADE_DATE > TRADE_DEADLINE:
            reasons.append("2026-27 regular-season trade deadline has passed")

        contract_type = str(row.get("contract_type", "") or "").lower()
        signed_year = pd.to_numeric(row.get("signed_year"), errors="coerce")
        signed = parse_date(row.get("contract_signed_date"))
        last_trade = parse_date(row.get("latest_trade_date"))
        event_type = str(row.get("contract_event_type", "") or "").lower()

        # Current-year free-agent contracts: later of Dec. 15 or three months
        # from signing. Missing dates cannot certify eligibility.
        if contract_type.startswith("free agent") and pd.notna(signed_year) and int(signed_year) == MIN_CONTRACT_SEASON:
            earliest = date(MIN_CONTRACT_SEASON, 12, 15)
            if signed:
                earliest = max(earliest, signed + relativedelta(months=3))
            elif TRADE_DATE >= earliest:
                unknowns.append("Free-agent signing date missing: three-month rule unresolved")
            if TRADE_DATE < earliest:
                reasons.append(f"Free-agent signing restriction until {earliest}")

        # Newly signed rookie and two-way contracts: 30-day waiting period.
        if pd.notna(signed_year) and int(signed_year) == MIN_CONTRACT_SEASON:
            if contract_type.strip() == "rookie" or "two-way" in contract_type:
                if signed is None:
                    unknowns.append("Rookie/two-way signing date missing: 30-day rule unresolved")
                elif TRADE_DATE < signed + timedelta(days=30):
                    reasons.append(f"30-day signing restriction until {signed + timedelta(days=30)}")

        # Some qualifying extensions impose a waiting period; do not assume
        # every extension does. Explicit designation needs corroborating data.
        if "designated veteran" in contract_type and pd.notna(signed_year) and int(signed_year) == TRADE_DATE.year:
            reasons.append("Designated-veteran one-year restriction (current signing year)")
        elif event_type == "extension" and signed and signed <= TRADE_DATE:
            unknowns.append("Extension restrictions may apply; extension type not reliably identified")

        # A recently traded player is NOT automatically forbidden from being
        # traded again alone. The two-month rule concerns aggregation and is
        # deliberately not applied to this individual-player filter.
        if last_trade is not None and last_trade > TRADE_DATE:
            unknowns.append("Recorded latest trade date is in the future; verify transaction data")

        # Consent, matched offer sheets, and special restrictions cannot be
        # verified reliably from the existing collection alone.
        # They remain unresolved rather than incorrectly marked legal.
        if not contract_type or contract_type in {"nan", "<na>"}:
            unknowns.append("Contract type missing")
        return not reasons, reasons, unknowns

    current = contracts_df[contracts_df["season_start"] == MIN_CONTRACT_SEASON].copy()
    contract_lookup = {}
    for _, row in current.iterrows():
        contract_lookup[(str(row["team"]), normalize_name(row["Name"]).casefold())] = row

    def filter_pool(pool, pool_name):
        remaining = {}
        unresolved = []
        removed = []
        for team, names in pool.items():
            accepted = []
            for player in names:
                record = contract_lookup.get((str(team), normalize_name(player).casefold()))
                if record is None:
                    removed.append((team, player, "No 2026-27 contract record for team"))
                    continue
                ok, reasons, unknowns = eligibility_record(record)
                if not ok:
                    removed.append((team, player, "; ".join(reasons)))
                    continue
                accepted.append(player)
                if unknowns:
                    unresolved.append((team, player, "; ".join(unknowns)))
            if accepted:
                remaining[team] = sorted(set(accepted))
        print(f"\n{pool_name}: {sum(map(len, remaining.values()))} players retained, {len(removed)} excluded")
        for team, player, reason in removed:
            print(f"  EXCLUDED | {team} | {player}: {reason}")
        for team, player, issue in unresolved:
            print(f"  UNVERIFIED | {team} | {player}: {issue}")
        return remaining

    # Filter ONLY the four user-selected GPT pools. No salary logic.
    eligible_required_targets = filter_pool(required_targets, "ELIGIBLE REQUIRED TARGETS")
    eligible_tradeable_targets = filter_pool(tradeable_targets, "ELIGIBLE TRADEABLE TARGETS")
    eligible_wanted_to_trade_away = filter_pool(wanted_to_trade_away, "ELIGIBLE WANTED TO TRADE AWAY")
    eligible_can_trade_away = filter_pool(can_trade_away, "ELIGIBLE CAN TRADE AWAY")

    print("\n" + "=" * 72)
    print("FINAL PLAYER DICTIONARIES — AFTER INDIVIDUAL CBA FILTERING")
    print("=" * 72)
    print("\nREQUIRED TARGETS:", eligible_required_targets)
    print("\nTRADEABLE TARGETS:", eligible_tradeable_targets)
    print("\nWANTED TO TRADE AWAY:", eligible_wanted_to_trade_away)
    print("\nCAN TRADE AWAY:", eligible_can_trade_away)

    # ============================================================
    # STAGES 6–8 — DISABLED FOR NOW
    # Trade generation, salary feasibility pruning, complete CBA checks,
    # and final legal-trade results will be re-enabled in a later version.
    # See commented reference implementation below.
    # ============================================================
    # ---- BEGIN DISABLED LEGACY TRADE GENERATOR / CBA VALIDATOR ----
    # # Viable Trades
    #
    # Money = float
    # PlayerKey = Tuple[str, str]
    #
    #
    # # ---------------------------------------------------------------------------
    # # 2026-27 system values from the supplied CBA reference.
    # # The salary-matching tier values are described as approximate in the source,
    # # so they live in SeasonRules rather than being buried in validation code.
    # # ---------------------------------------------------------------------------
    #
    #
    # @dataclass(frozen=True)
    # class SeasonRules:
    #     season_start: int = 2026
    #     salary_cap: Money = 164_961_000
    #     luxury_tax: Money = 200_428_000
    #     first_apron: Money = 209_015_000
    #     second_apron: Money = 221_686_000
    #
    #     matching_cushion: Money = 250_000
    #     expanded_low_threshold: Money = 8_846_000
    #     expanded_middle_threshold: Money = 35_383_000
    #     expanded_middle_addon: Money = 9_096_000
    #
    #     annual_cash_limit: Money = 8_495_000
    #
    #     # Rule 30: intentionally season-level/configurable, not hardcoded here.
    #     trade_deadline: Optional[date] = None
    #
    #     # Rule 18: sign-and-trade must occur before regular season begins.
    #     regular_season_start: Optional[date] = None
    #
    #     # Set True only when validating the post-deadline/postseason trade window.
    #     postseason_trade_window_open: bool = False
    #
    #     # For 2026-27, the first future draft is normally 2027.
    #     first_future_draft_year: int = 2027
    #
    #
    # @dataclass(frozen=True)
    # class PlayerData:
    #     name: str
    #     team: str
    #     team_abbreviation: Optional[str]
    #     salary: Money
    #     season_start: int
    #     player_option: bool
    #     team_option: bool
    #     qualifying_offer: bool
    #     extension_eligible: bool
    #     guarantee_deadline: bool
    #     signed_year: Optional[int]
    #     contract_type: str
    #     contract_start: Optional[int]
    #     contract_end: Optional[int]
    #     contract_years: Optional[int]
    #     contract_value: Optional[Money]
    #     guaranteed_at_signing: Optional[Money]
    #     practical_guaranteed: Optional[Money]
    #
    #
    # @dataclass(frozen=True)
    # class TeamData:
    #     team: str
    #     season_start: int
    #     salary_cap: Money
    #     active_cap: Money
    #     total_cap_allocations: Money
    #     cap_space: Money
    #     first_apron: Money
    #     first_apron_allocations: Money
    #     first_apron_space: Money
    #     second_apron: Money
    #     second_apron_allocations: Money
    #     second_apron_space: Money
    #     above_cap: bool
    #     above_first_apron: bool
    #     above_second_apron: bool
    #
    #
    # @dataclass(frozen=True)
    # class DraftAsset:
    #     team: str
    #     draft_year: int
    #     round: int
    #     direction: str
    #     asset: str
    #     details: str
    #     is_swap: bool
    #     is_protected: bool
    #     is_conditional: bool
    #
    #
    # @dataclass(frozen=True)
    # class TradeException:
    #     amount: Money
    #     expires: Optional[date] = None
    #     hard_caps_first_apron: bool = False
    #
    #
    # @dataclass(frozen=True)
    # class Violation:
    #     rule: str
    #     code: str
    #     message: str
    #     team: Optional[str] = None
    #     player: Optional[str] = None
    #
    #
    # @dataclass
    # class ValidationResult:
    #     passes: bool
    #     violations: List[Violation] = field(default_factory=list)
    #     undetermined: List[str] = field(default_factory=list)
    #     mechanisms: Dict[str, str] = field(default_factory=dict)
    #     team_salary_detail: Dict[str, Dict[str, Money]] = field(default_factory=dict)
    #
    #     @property
    #     def legal(self) -> bool:
    #         """Alias used by callers that prefer `result.legal`."""
    #         return self.passes
    #
    #
    # @dataclass(frozen=True)
    # class RejectedTrade:
    #     index: int
    #     trade: Mapping[str, Any]
    #     result: ValidationResult
    #
    #
    # @dataclass(frozen=True)
    # class TeamFlow:
    #     outgoing_players: Tuple[Tuple[str, str], ...]  # (player, destination)
    #     incoming_players: Tuple[Tuple[str, str], ...]  # (player, origin)
    #     cash_sent: Money
    #     cash_received: Money
    #     outgoing_picks: Tuple[Mapping[str, Any], ...]
    #     incoming_picks: Tuple[Mapping[str, Any], ...]
    #
    #
    # # ---------------------------------------------------------------------------
    # # Small helpers
    # # ---------------------------------------------------------------------------
    #
    #
    # def _as_money(value: Any, default: Money = 0.0) -> Money:
    #     if value in (None, ""):
    #         return default
    #     if isinstance(value, bool):
    #         return float(value)
    #     if isinstance(value, (int, float)):
    #         return float(value)
    #     text = str(value).strip().replace("$", "").replace(",", "")
    #     if not text:
    #         return default
    #     return float(text)
    #
    #
    # def _as_int(value: Any) -> Optional[int]:
    #     if value in (None, ""):
    #         return None
    #     return int(value)
    #
    #
    # def _as_bool(value: Any) -> bool:
    #     if isinstance(value, bool):
    #         return value
    #     if value in (None, ""):
    #         return False
    #     if isinstance(value, (int, float)):
    #         return bool(value)
    #     return str(value).strip().lower() in {"1", "true", "yes", "y"}
    #
    #
    # def _as_date(value: Any) -> Optional[date]:
    #     if value in (None, ""):
    #         return None
    #     if isinstance(value, datetime):
    #         return value.date()
    #     if isinstance(value, date):
    #         return value
    #     if isinstance(value, str):
    #         return date.fromisoformat(value.strip())
    #     raise TypeError(f"Cannot convert {value!r} to date.")
    #
    #
    # def _add_months(d: date, months: int) -> date:
    #     month_index = (d.month - 1) + months
    #     year = d.year + month_index // 12
    #     month = month_index % 12 + 1
    #     day = min(d.day, calendar.monthrange(year, month)[1])
    #     return date(year, month, day)
    #
    #
    # def _add_year(d: date) -> date:
    #     try:
    #         return d.replace(year=d.year + 1)
    #     except ValueError:
    #         # Feb. 29 -> Feb. 28 in a non-leap year.
    #         return d.replace(month=2, day=28, year=d.year + 1)
    #
    #
    # def _first_value(row: Mapping[str, Any], *names: str) -> Any:
    #     for name in names:
    #         if name in row and row[name] not in (None, ""):
    #             return row[name]
    #     return None
    #
    #
    # def _normalize_header(value: Any) -> str:
    #     if value is None:
    #         return ""
    #     return str(value).strip().lower().replace(" ", "_").replace("/", "_")
    #
    #
    # def _rows_to_dicts(headers: Sequence[Any], rows: Sequence[Sequence[Any]]) -> List[Dict[str, Any]]:
    #     keys = [_normalize_header(h) for h in headers]
    #     output: List[Dict[str, Any]] = []
    #     for row in rows:
    #         if not any(value not in (None, "") for value in row):
    #             continue
    #         item = {
    #             key: row[i] if i < len(row) else None
    #             for i, key in enumerate(keys)
    #             if key
    #         }
    #         output.append(item)
    #     return output
    #
    #
    # # ---------------------------------------------------------------------------
    # # Workbook loader
    # # ---------------------------------------------------------------------------
    #
    #
    # class EngineData:
    #     """Normalized view of the uploaded NBA trade-engine workbook."""
    #
    #     SECTION_HEADINGS = {
    #         "CONTRACTS",
    #         "TRANSACTIONS",
    #         "TEAM CAP / APRONS",
    #         "TRADE EXCEPTIONS",
    #         "CONTRACT DEADLINES",
    #         "DRAFT ASSETS",
    #     }
    #
    #     def __init__(self, season_start: int):
    #         self.season_start = season_start
    #         self.players: Dict[PlayerKey, PlayerData] = {}
    #         self.teams: Dict[str, TeamData] = {}
    #         self.draft_assets: Dict[str, List[DraftAsset]] = {}
    #         self.trade_exceptions: Dict[str, List[TradeException]] = {}
    #
    #     @classmethod
    #     def from_xlsx(cls, path: str | Path, season_start: int = 2026) -> "EngineData":
    #         path = Path(path)
    #         wb = load_workbook(path, read_only=True, data_only=True)
    #         data = cls(season_start=season_start)
    #
    #         try:
    #             for ws in wb.worksheets:
    #                 rows = [tuple(row) for row in ws.iter_rows(values_only=True)]
    #                 data._parse_team_sheet(ws.title, rows)
    #         finally:
    #             wb.close()
    #
    #         return data
    #
    #     def _parse_team_sheet(self, sheet_name: str, rows: Sequence[Sequence[Any]]) -> None:
    #         sections: Dict[str, Tuple[int, int]] = {}
    #         heading_positions: List[Tuple[int, str]] = []
    #
    #         for i, row in enumerate(rows):
    #             first = row[0] if row else None
    #             if first in self.SECTION_HEADINGS:
    #                 heading_positions.append((i, str(first)))
    #
    #         for pos, (start, heading) in enumerate(heading_positions):
    #             end = heading_positions[pos + 1][0] if pos + 1 < len(heading_positions) else len(rows)
    #             sections[heading] = (start, end)
    #
    #         if "CONTRACTS" in sections:
    #             self._parse_contracts(sheet_name, rows, *sections["CONTRACTS"])
    #         if "TEAM CAP / APRONS" in sections:
    #             self._parse_cap(sheet_name, rows, *sections["TEAM CAP / APRONS"])
    #         if "TRADE EXCEPTIONS" in sections:
    #             self._parse_trade_exceptions(sheet_name, rows, *sections["TRADE EXCEPTIONS"])
    #         if "DRAFT ASSETS" in sections:
    #             self._parse_draft_assets(sheet_name, rows, *sections["DRAFT ASSETS"])
    #
    #     def _section_dicts(
    #         self,
    #         rows: Sequence[Sequence[Any]],
    #         start: int,
    #         end: int,
    #     ) -> List[Dict[str, Any]]:
    #         if start + 1 >= end:
    #             return []
    #         headers = rows[start + 1]
    #         if not headers or headers[0] in (None, "Info"):
    #             return []
    #         return _rows_to_dicts(headers, rows[start + 2 : end])
    #
    #     def _parse_contracts(
    #         self,
    #         sheet_name: str,
    #         rows: Sequence[Sequence[Any]],
    #         start: int,
    #         end: int,
    #     ) -> None:
    #         for row in self._section_dicts(rows, start, end):
    #             if _as_int(row.get("season_start")) != self.season_start:
    #                 continue
    #             name = str(row.get("name") or "").strip()
    #             team = str(row.get("team") or sheet_name).strip()
    #             if not name:
    #                 continue
    #
    #             self.players[(team, name)] = PlayerData(
    #                 name=name,
    #                 team=team,
    #                 team_abbreviation=(str(row.get("team_abbreviation")) if row.get("team_abbreviation") else None),
    #                 salary=_as_money(row.get("salary")),
    #                 season_start=self.season_start,
    #                 player_option=_as_bool(row.get("player_option")),
    #                 team_option=_as_bool(row.get("team_option")),
    #                 qualifying_offer=_as_bool(row.get("qualifying_offer")),
    #                 extension_eligible=_as_bool(row.get("extension_eligible")),
    #                 guarantee_deadline=_as_bool(row.get("guarantee_deadline")),
    #                 signed_year=_as_int(row.get("signed_year")),
    #                 contract_type=str(row.get("contract_type") or ""),
    #                 contract_start=_as_int(row.get("contract_start")),
    #                 contract_end=_as_int(row.get("contract_end")),
    #                 contract_years=_as_int(row.get("contract_years")),
    #                 contract_value=(None if row.get("contract_value") in (None, "") else _as_money(row.get("contract_value"))),
    #                 guaranteed_at_signing=(None if row.get("guaranteed_at_signing") in (None, "") else _as_money(row.get("guaranteed_at_signing"))),
    #                 practical_guaranteed=(None if row.get("practical_guaranteed") in (None, "") else _as_money(row.get("practical_guaranteed"))),
    #             )
    #
    #     def _parse_cap(
    #         self,
    #         sheet_name: str,
    #         rows: Sequence[Sequence[Any]],
    #         start: int,
    #         end: int,
    #     ) -> None:
    #         for row in self._section_dicts(rows, start, end):
    #             if _as_int(row.get("season_start")) != self.season_start:
    #                 continue
    #             team = str(row.get("team") or sheet_name).strip()
    #             self.teams[team] = TeamData(
    #                 team=team,
    #                 season_start=self.season_start,
    #                 salary_cap=_as_money(row.get("salary_cap")),
    #                 active_cap=_as_money(row.get("active_cap")),
    #                 total_cap_allocations=_as_money(row.get("total_cap_allocations")),
    #                 cap_space=_as_money(row.get("cap_space")),
    #                 first_apron=_as_money(row.get("first_apron")),
    #                 first_apron_allocations=_as_money(row.get("first_apron_allocations")),
    #                 first_apron_space=_as_money(row.get("first_apron_space")),
    #                 second_apron=_as_money(row.get("second_apron")),
    #                 second_apron_allocations=_as_money(row.get("second_apron_allocations")),
    #                 second_apron_space=_as_money(row.get("second_apron_space")),
    #                 above_cap=_as_bool(row.get("above_cap")),
    #                 above_first_apron=_as_bool(row.get("above_first_apron")),
    #                 above_second_apron=_as_bool(row.get("above_second_apron")),
    #             )
    #             break
    #
    #     def _parse_trade_exceptions(
    #         self,
    #         sheet_name: str,
    #         rows: Sequence[Sequence[Any]],
    #         start: int,
    #         end: int,
    #     ) -> None:
    #         # The supplied workbook currently has "No data found" here. This
    #         # parser is intentionally schema-tolerant for future workbook versions.
    #         parsed = self._section_dicts(rows, start, end)
    #         exceptions: List[TradeException] = []
    #
    #         for row in parsed:
    #             amount_raw = _first_value(
    #                 row,
    #                 "remaining_amount",
    #                 "amount_remaining",
    #                 "amount",
    #                 "trade_exception_amount",
    #                 "tpe_amount",
    #                 "value",
    #             )
    #             if amount_raw in (None, ""):
    #                 continue
    #
    #             expires_raw = _first_value(
    #                 row,
    #                 "expiration_date",
    #                 "expires",
    #                 "expiry_date",
    #             )
    #             expires = None
    #             if isinstance(expires_raw, datetime):
    #                 expires = expires_raw.date()
    #             elif isinstance(expires_raw, date):
    #                 expires = expires_raw
    #             elif isinstance(expires_raw, str) and expires_raw.strip():
    #                 try:
    #                     expires = date.fromisoformat(expires_raw.strip())
    #                 except ValueError:
    #                     expires = None
    #
    #             exceptions.append(
    #                 TradeException(
    #                     amount=_as_money(amount_raw),
    #                     expires=expires,
    #                     hard_caps_first_apron=_as_bool(
    #                         _first_value(row, "hard_caps_first_apron", "first_apron_hard_cap")
    #                     ),
    #                 )
    #             )
    #
    #         if exceptions:
    #             self.trade_exceptions[sheet_name] = exceptions
    #
    #     def _parse_draft_assets(
    #         self,
    #         sheet_name: str,
    #         rows: Sequence[Sequence[Any]],
    #         start: int,
    #         end: int,
    #     ) -> None:
    #         assets: List[DraftAsset] = []
    #         for row in self._section_dicts(rows, start, end):
    #             year = _as_int(row.get("draft_year"))
    #             rnd = _as_int(row.get("round"))
    #             if year is None or rnd is None:
    #                 continue
    #             assets.append(
    #                 DraftAsset(
    #                     team=str(row.get("team") or sheet_name),
    #                     draft_year=year,
    #                     round=rnd,
    #                     direction=str(row.get("direction") or "").upper(),
    #                     asset=str(row.get("asset") or ""),
    #                     details=str(row.get("details") or ""),
    #                     is_swap=_as_bool(row.get("is_swap")),
    #                     is_protected=_as_bool(row.get("is_protected")),
    #                     is_conditional=_as_bool(row.get("is_conditional")),
    #                 )
    #             )
    #         if assets:
    #             self.draft_assets[sheet_name] = assets
    #
    #
    # # ---------------------------------------------------------------------------
    # # Validator
    # # ---------------------------------------------------------------------------
    #
    #
    # class CBAValidator:
    #     def __init__(
    #         self,
    #         data: EngineData,
    #         trade_date: date,
    #         rules: Optional[SeasonRules] = None,
    #         player_overrides: Optional[Mapping[PlayerKey, Mapping[str, Any]]] = None,
    #         team_overrides: Optional[Mapping[str, Mapping[str, Any]]] = None,
    #     ):
    #         self.data = data
    #         self.trade_date = trade_date
    #         self.rules = rules or SeasonRules(season_start=data.season_start)
    #         self.player_overrides: Mapping[PlayerKey, Mapping[str, Any]] = player_overrides or {}
    #         self.team_overrides: Mapping[str, Mapping[str, Any]] = team_overrides or {}
    #
    #     @classmethod
    #     def from_xlsx(
    #         cls,
    #         path: str | Path,
    #         trade_date: date,
    #         rules: Optional[SeasonRules] = None,
    #         player_overrides: Optional[Mapping[PlayerKey, Mapping[str, Any]]] = None,
    #         team_overrides: Optional[Mapping[str, Mapping[str, Any]]] = None,
    #     ) -> "CBAValidator":
    #         resolved_rules = rules or SeasonRules()
    #         data = EngineData.from_xlsx(path, season_start=resolved_rules.season_start)
    #         return cls(
    #             data=data,
    #             trade_date=trade_date,
    #             rules=resolved_rules,
    #             player_overrides=player_overrides,
    #             team_overrides=team_overrides,
    #         )
    #
    #     # ------------------------- public API ---------------------------------
    #
    #     def validate_trade(self, trade: Mapping[str, Any]) -> ValidationResult:
    #         violations: List[Violation] = []
    #         undetermined: List[str] = []
    #         mechanisms: Dict[str, str] = {}
    #         team_salary_detail: Dict[str, Dict[str, Money]] = {}
    #
    #         flows = self._build_team_flows(trade, violations)
    #         if not flows:
    #             if not violations:
    #                 violations.append(
    #                     Violation(
    #                         rule="Structure",
    #                         code="NO_PARTICIPANTS",
    #                         message="Trade has no participating teams/moves.",
    #                     )
    #                 )
    #             return ValidationResult(False, violations, undetermined, mechanisms, team_salary_detail)
    #
    #         self._validate_trade_deadline(violations)
    #         self._validate_player_ownership_and_eligibility(flows, violations, undetermined)
    #
    #         # Only perform salary math for teams whose player references are known.
    #         self._validate_each_team_salary(
    #             flows,
    #             violations,
    #             undetermined,
    #             mechanisms,
    #             team_salary_detail,
    #         )
    #
    #         self._validate_cash(flows, violations, undetermined, team_salary_detail)
    #         self._validate_draft_assets(flows, violations, undetermined)
    #
    #         return ValidationResult(
    #             passes=not violations,
    #             violations=violations,
    #             undetermined=undetermined,
    #             mechanisms=mechanisms,
    #             team_salary_detail=team_salary_detail,
    #         )
    #
    #     def is_trade_legal(self, trade: Mapping[str, Any]) -> bool:
    #         return self.validate_trade(trade).passes
    #
    #     def iter_legal_trades(
    #         self,
    #         trades: Iterable[Mapping[str, Any]],
    #     ) -> Iterator[Mapping[str, Any]]:
    #         """Memory-friendly filter for generic trade-variation iterables."""
    #         for trade in trades:
    #             if self.validate_trade(trade).passes:
    #                 yield trade
    #
    #     def filter_trades(
    #         self,
    #         trades: Iterable[Mapping[str, Any]],
    #         *,
    #         keep_rejections: bool = False,
    #         rejection_limit: Optional[int] = None,
    #     ) -> Tuple[List[Mapping[str, Any]], List[RejectedTrade]]:
    #         legal: List[Mapping[str, Any]] = []
    #         rejected: List[RejectedTrade] = []
    #
    #         for index, trade in enumerate(trades):
    #             result = self.validate_trade(trade)
    #             if result.passes:
    #                 legal.append(trade)
    #             elif keep_rejections and (
    #                 rejection_limit is None or len(rejected) < rejection_limit
    #             ):
    #                 rejected.append(RejectedTrade(index=index, trade=trade, result=result))
    #
    #         return legal, rejected
    #
    #     # ------------------------- trade normalization ------------------------
    #
    #     def _build_team_flows(
    #         self,
    #         trade: Mapping[str, Any],
    #         violations: List[Violation],
    #     ) -> Dict[str, TeamFlow]:
    #         outgoing_players: Dict[str, List[Tuple[str, str]]] = {}
    #         incoming_players: Dict[str, List[Tuple[str, str]]] = {}
    #         outgoing_picks: Dict[str, List[Mapping[str, Any]]] = {}
    #         incoming_picks: Dict[str, List[Mapping[str, Any]]] = {}
    #         cash_sent: Dict[str, Money] = {}
    #         cash_received: Dict[str, Money] = {}
    #
    #         seen_player_moves: set[PlayerKey] = set()
    #
    #         moves = trade.get("moves", ())
    #         if not isinstance(moves, (tuple, list)):
    #             violations.append(
    #                 Violation(
    #                     rule="Structure",
    #                     code="MOVES_NOT_SEQUENCE",
    #                     message="trade['moves'] must be a list or tuple.",
    #                 )
    #             )
    #             return {}
    #
    #         for move_index, move in enumerate(moves):
    #             if not isinstance(move, Mapping):
    #                 violations.append(
    #                     Violation(
    #                         rule="Structure",
    #                         code="MOVE_NOT_MAPPING",
    #                         message=f"Move #{move_index} is not a mapping.",
    #                     )
    #                 )
    #                 continue
    #
    #             from_team = str(move.get("from_team") or "").strip()
    #             to_team = str(move.get("to_team") or "").strip()
    #             if not from_team or not to_team or from_team == to_team:
    #                 violations.append(
    #                     Violation(
    #                         rule="Structure",
    #                         code="INVALID_MOVE_TEAMS",
    #                         message=f"Move #{move_index} has invalid from/to teams.",
    #                         team=from_team or to_team or None,
    #                     )
    #                 )
    #                 continue
    #
    #             outgoing_players.setdefault(from_team, [])
    #             incoming_players.setdefault(from_team, [])
    #             outgoing_players.setdefault(to_team, [])
    #             incoming_players.setdefault(to_team, [])
    #             outgoing_picks.setdefault(from_team, [])
    #             incoming_picks.setdefault(from_team, [])
    #             outgoing_picks.setdefault(to_team, [])
    #             incoming_picks.setdefault(to_team, [])
    #             cash_sent.setdefault(from_team, 0.0)
    #             cash_received.setdefault(from_team, 0.0)
    #             cash_sent.setdefault(to_team, 0.0)
    #             cash_received.setdefault(to_team, 0.0)
    #
    #             players = move.get("players", ()) or ()
    #             for player_raw in players:
    #                 player = str(player_raw).strip()
    #                 key = (from_team, player)
    #                 if key in seen_player_moves:
    #                     violations.append(
    #                         Violation(
    #                             rule="Structure",
    #                             code="PLAYER_MOVED_TWICE",
    #                             message=f"{player} is sent more than once by {from_team}.",
    #                             team=from_team,
    #                             player=player,
    #                         )
    #                     )
    #                     continue
    #                 seen_player_moves.add(key)
    #                 outgoing_players[from_team].append((player, to_team))
    #                 incoming_players[to_team].append((player, from_team))
    #
    #             picks = move.get("picks", ()) or ()
    #             for pick_raw in picks:
    #                 if not isinstance(pick_raw, Mapping):
    #                     violations.append(
    #                         Violation(
    #                             rule="25-29",
    #                             code="INVALID_PICK_OBJECT",
    #                             message="Each move['picks'] entry must be a mapping.",
    #                             team=from_team,
    #                         )
    #                     )
    #                     continue
    #                 pick = dict(pick_raw)
    #                 outgoing_picks[from_team].append(pick)
    #                 incoming_picks[to_team].append(pick)
    #
    #             cash = _as_money(move.get("cash"), 0.0)
    #             if cash < 0:
    #                 violations.append(
    #                     Violation(
    #                         rule="24",
    #                         code="NEGATIVE_CASH",
    #                         message="Cash consideration cannot be negative.",
    #                         team=from_team,
    #                     )
    #                 )
    #             elif cash:
    #                 cash_sent[from_team] += cash
    #                 cash_received[to_team] += cash
    #
    #         teams = set(outgoing_players) | set(incoming_players)
    #         flows: Dict[str, TeamFlow] = {}
    #         for team in teams:
    #             flows[team] = TeamFlow(
    #                 outgoing_players=tuple(outgoing_players.get(team, ())),
    #                 incoming_players=tuple(incoming_players.get(team, ())),
    #                 cash_sent=cash_sent.get(team, 0.0),
    #                 cash_received=cash_received.get(team, 0.0),
    #                 outgoing_picks=tuple(outgoing_picks.get(team, ())),
    #                 incoming_picks=tuple(incoming_picks.get(team, ())),
    #             )
    #
    #         # If trade['teams'] is supplied, ensure every declared team participates.
    #         declared = trade.get("teams")
    #         if declared:
    #             declared_set = {str(t) for t in declared}
    #             actual_set = set(flows)
    #             if declared_set != actual_set:
    #                 violations.append(
    #                     Violation(
    #                         rule="31",
    #                         code="DECLARED_TEAM_MISMATCH",
    #                         message=(
    #                             f"Declared teams {sorted(declared_set)} do not match teams "
    #                             f"actually present in moves {sorted(actual_set)}."
    #                         ),
    #                     )
    #                 )
    #
    #         return flows
    #
    #     # ------------------------- rule helpers -------------------------------
    #
    #     def _player_override(self, team: str, player: str) -> Mapping[str, Any]:
    #         return self.player_overrides.get((team, player), {})
    #
    #     def _team_override(self, team: str) -> Mapping[str, Any]:
    #         return self.team_overrides.get(team, {})
    #
    #     def _player_data(self, team: str, player: str) -> Optional[PlayerData]:
    #         return self.data.players.get((team, player))
    #
    #     def _team_data(self, team: str) -> Optional[TeamData]:
    #         return self.data.teams.get(team)
    #
    #     def _outgoing_trade_salary(self, origin: str, player: str) -> Money:
    #         override = self._player_override(origin, player)
    #         if override.get("outgoing_trade_salary") not in (None, ""):
    #             return _as_money(override["outgoing_trade_salary"])
    #         p = self._player_data(origin, player)
    #         return p.salary if p else 0.0
    #
    #     def _incoming_trade_salary(self, origin: str, player: str) -> Money:
    #         override = self._player_override(origin, player)
    #         if override.get("incoming_trade_salary") not in (None, ""):
    #             return _as_money(override["incoming_trade_salary"])
    #         p = self._player_data(origin, player)
    #         return p.salary if p else 0.0
    #
    #     def _apron_salary(self, origin: str, player: str) -> Money:
    #         override = self._player_override(origin, player)
    #         if override.get("apron_salary") not in (None, ""):
    #             return _as_money(override["apron_salary"])
    #         # A trade kicker or special cap treatment can be represented by
    #         # overriding apron_salary. Otherwise current salary is the best field
    #         # present in the supplied workbook.
    #         p = self._player_data(origin, player)
    #         return p.salary if p else 0.0
    #
    #     # ------------------------- Rule 30 ------------------------------------
    #
    #     def _validate_trade_deadline(self, violations: List[Violation]) -> None:
    #         deadline = self.rules.trade_deadline
    #         if deadline is None:
    #             return
    #
    #         if self.trade_date > deadline and not self.rules.postseason_trade_window_open:
    #             violations.append(
    #                 Violation(
    #                     rule="30",
    #                     code="TRADE_DEADLINE_CLOSED",
    #                     message=(
    #                         f"Trade date {self.trade_date.isoformat()} is after the configured "
    #                         f"trade deadline {deadline.isoformat()} and the postseason trade "
    #                         "window is not marked open."
    #                     ),
    #                 )
    #             )
    #
    #     # ------------------------- Rules 7-23 player restrictions -------------
    #
    #     def _validate_player_ownership_and_eligibility(
    #         self,
    #         flows: Mapping[str, TeamFlow],
    #         violations: List[Violation],
    #         undetermined: List[str],
    #     ) -> None:
    #         for origin, flow in flows.items():
    #             outgoing_count = len(flow.outgoing_players)
    #
    #             for player, destination in flow.outgoing_players:
    #                 pdata = self._player_data(origin, player)
    #                 if pdata is None:
    #                     violations.append(
    #                         Violation(
    #                             rule="Structure",
    #                             code="PLAYER_NOT_ON_ORIGIN_TEAM",
    #                             message=(
    #                                 f"{player} does not have a {self.rules.season_start}-{str(self.rules.season_start + 1)[-2:]} "
    #                                 f"contract row for {origin} in the engine workbook."
    #                             ),
    #                             team=origin,
    #                             player=player,
    #                         )
    #                     )
    #                     continue
    #
    #                 ov = self._player_override(origin, player)
    #                 ctype = pdata.contract_type.lower()
    #
    #                 # Rule 7: newly signed standard free agent restriction.
    #                 # Do not apply this to the CURRENT sign-and-trade itself; Rule 18
    #                 # specifically governs that transaction. Historical S&T contract
    #                 # types are not automatically treated as a current S&T event.
    #                 current_sign_and_trade = _as_bool(ov.get("being_signed_and_traded"))
    #                 if (
    #                     ctype.startswith("free agent")
    #                     and pdata.signed_year == self.rules.season_start
    #                     and not current_sign_and_trade
    #                 ):
    #                     signed_date = _as_date(ov.get("signed_date"))
    #                     dec15 = date(self.rules.season_start, 12, 15)
    #
    #                     if signed_date is not None:
    #                         eligible = max(_add_months(signed_date, 3), dec15)
    #                         if self.trade_date < eligible:
    #                             violations.append(
    #                                 Violation(
    #                                     rule="7",
    #                                     code="NEW_FREE_AGENT_TRADE_RESTRICTION",
    #                                     message=(
    #                                         f"{player} is a current-cap-year free-agent signing and "
    #                                         f"cannot be traded before {eligible.isoformat()}."
    #                                     ),
    #                                     team=origin,
    #                                     player=player,
    #                                 )
    #                             )
    #                     elif self.trade_date < dec15:
    #                         # Even without the exact signing date, the later-of rule
    #                         # proves the player is ineligible before Dec. 15.
    #                         violations.append(
    #                             Violation(
    #                                 rule="7",
    #                                 code="NEW_FREE_AGENT_BEFORE_DEC15",
    #                                 message=(
    #                                     f"{player} signed as a free agent in {self.rules.season_start} "
    #                                     f"and the trade date is before {dec15.isoformat()}."
    #                                 ),
    #                                 team=origin,
    #                                 player=player,
    #                             )
    #                         )
    #                     else:
    #                         undetermined.append(
    #                             f"Rule 7: exact signed_date missing for {origin} / {player}; "
    #                             "three-month portion of the restriction cannot be proven."
    #                         )
    #
    #                 # Rule 8: January 15 restriction applies only when its factual
    #                 # predicates are satisfied. Workbook does not identify them.
    #                 if _as_bool(ov.get("jan15_restricted")):
    #                     signed_date = _as_date(ov.get("signed_date"))
    #                     jan15 = date(self.rules.season_start + 1, 1, 15)
    #                     if signed_date is None:
    #                         eligible = jan15
    #                         undetermined.append(
    #                             f"Rule 8: signed_date missing for Jan. 15 restricted player "
    #                             f"{origin} / {player}; using Jan. 15 only is not sufficient to certify."
    #                         )
    #                     else:
    #                         eligible = max(_add_months(signed_date, 3), jan15)
    #                     if self.trade_date < eligible:
    #                         violations.append(
    #                             Violation(
    #                                 rule="8",
    #                                 code="JAN15_RESTRICTION",
    #                                 message=f"{player} is not trade eligible until {eligible.isoformat()}.",
    #                                 team=origin,
    #                                 player=player,
    #                             )
    #                         )
    #
    #                 # Rule 9: newly signed draft pick, 30 days.
    #                 is_draft_rookie = ov.get("is_draft_rookie")
    #                 if is_draft_rookie is None:
    #                     is_draft_rookie = pdata.contract_type.strip().lower() == "rookie"
    #                 if _as_bool(is_draft_rookie) and pdata.signed_year == self.rules.season_start:
    #                     signed_date = _as_date(ov.get("signed_date"))
    #                     if signed_date is None:
    #                         undetermined.append(
    #                             f"Rule 9: signed_date missing for draft rookie {origin} / {player}."
    #                         )
    #                     elif self.trade_date < signed_date + timedelta(days=30):
    #                         violations.append(
    #                             Violation(
    #                                 rule="9",
    #                                 code="ROOKIE_30_DAY_RESTRICTION",
    #                                 message=f"{player} is within 30 days of signing his rookie contract.",
    #                                 team=origin,
    #                                 player=player,
    #                             )
    #                         )
    #
    #                 # Rule 10: Two-Way player, 30 days.
    #                 is_two_way = ov.get("is_two_way")
    #                 if is_two_way is None:
    #                     is_two_way = "two-way" in ctype
    #                 if _as_bool(is_two_way) and pdata.signed_year == self.rules.season_start:
    #                     signed_date = _as_date(ov.get("signed_date"))
    #                     if signed_date is None:
    #                         undetermined.append(
    #                             f"Rule 10: signed_date missing for Two-Way player {origin} / {player}."
    #                         )
    #                     elif self.trade_date < signed_date + timedelta(days=30):
    #                         violations.append(
    #                             Violation(
    #                                 rule="10",
    #                                 code="TWO_WAY_30_DAY_RESTRICTION",
    #                                 message=f"{player} is within 30 days of signing a Two-Way contract.",
    #                                 team=origin,
    #                                 player=player,
    #                             )
    #                         )
    #
    #                 # Rule 11: recently acquired player cannot be aggregated for
    #                 # two months in applicable situations.
    #                 acquired_date = _as_date(ov.get("acquired_date"))
    #                 if acquired_date is not None and outgoing_count > 1:
    #                     eligible = _add_months(acquired_date, 2)
    #                     if self.trade_date < eligible:
    #                         violations.append(
    #                             Violation(
    #                                 rule="11",
    #                                 code="RECENTLY_ACQUIRED_AGGREGATION",
    #                                 message=(
    #                                     f"{player} was acquired on {acquired_date.isoformat()} and "
    #                                     "is being aggregated with another outgoing player before "
    #                                     f"{eligible.isoformat()}."
    #                                 ),
    #                                 team=origin,
    #                                 player=player,
    #                             )
    #                         )
    #
    #                 # Rule 12: one-year Bird-rights veto.
    #                 if _as_bool(ov.get("bird_trade_veto")) and not _as_bool(ov.get("trade_consent")):
    #                     violations.append(
    #                         Violation(
    #                             rule="12",
    #                             code="BIRD_RIGHTS_VETO_NO_CONSENT",
    #                             message=f"{player} has a Bird/Early Bird trade veto and consent is not recorded.",
    #                             team=origin,
    #                             player=player,
    #                         )
    #                     )
    #
    #                 # Rule 13: contractual no-trade clause.
    #                 if _as_bool(ov.get("no_trade_clause")) and not _as_bool(ov.get("trade_consent")):
    #                     violations.append(
    #                         Violation(
    #                             rule="13",
    #                             code="NO_TRADE_CLAUSE_NO_CONSENT",
    #                             message=f"{player} has a no-trade clause and consent is not recorded.",
    #                             team=origin,
    #                             player=player,
    #                         )
    #                     )
    #
    #                 # Rule 14: matched RFA offer sheet.
    #                 matched_date = _as_date(ov.get("matched_rfa_date"))
    #                 if matched_date is not None and self.trade_date < _add_year(matched_date):
    #                     if not _as_bool(ov.get("trade_consent")):
    #                         violations.append(
    #                             Violation(
    #                                 rule="14",
    #                                 code="MATCHED_RFA_NO_CONSENT",
    #                                 message=(
    #                                     f"{player} is within one year of a matched RFA offer sheet "
    #                                     "and consent is not recorded."
    #                                 ),
    #                                 team=origin,
    #                                 player=player,
    #                             )
    #                         )
    #                     offer_team = ov.get("matched_rfa_offer_team")
    #                     if offer_team and destination == offer_team:
    #                         violations.append(
    #                             Violation(
    #                                 rule="14",
    #                                 code="MATCHED_RFA_ORIGINAL_OFFER_TEAM",
    #                                 message=(
    #                                     f"{player} cannot be traded to original offer-sheet team "
    #                                     f"{destination} within one year of the match."
    #                                 ),
    #                                 team=origin,
    #                                 player=player,
    #                             )
    #                         )
    #
    #                 # Rule 15: designated veteran contract/extension, one year.
    #                 designated_date = _as_date(ov.get("designated_veteran_signed_date"))
    #                 is_designated = "designated veteran" in ctype
    #                 if designated_date is not None:
    #                     if self.trade_date < _add_year(designated_date):
    #                         violations.append(
    #                             Violation(
    #                                 rule="15",
    #                                 code="DESIGNATED_VETERAN_ONE_YEAR",
    #                                 message=(
    #                                     f"{player} is within one year of a qualifying Designated "
    #                                     "Veteran signing/extension."
    #                                 ),
    #                                 team=origin,
    #                                 player=player,
    #                             )
    #                         )
    #                 elif is_designated and pdata.signed_year == self.trade_date.year:
    #                     # Any already-signed contract in the same calendar year is
    #                     # necessarily less than one year old.
    #                     violations.append(
    #                         Violation(
    #                             rule="15",
    #                             code="DESIGNATED_VETERAN_SAME_YEAR",
    #                             message=(
    #                                 f"{player}'s workbook contract is a Designated Veteran deal "
    #                                 f"signed in {pdata.signed_year}; on {self.trade_date.isoformat()} "
    #                                 "one year cannot yet have elapsed."
    #                             ),
    #                             team=origin,
    #                             player=player,
    #                         )
    #                     )
    #
    #                 # Rule 16: certain extensions/renegotiations, six months.
    #                 for field_name, code, label in (
    #                     ("extension_signed_date", "EXTENSION_SIX_MONTH", "extension"),
    #                     ("renegotiated_date", "RENEGOTIATION_SIX_MONTH", "renegotiation"),
    #                 ):
    #                     event_date = _as_date(ov.get(field_name))
    #                     if event_date is not None and self.trade_date < _add_months(event_date, 6):
    #                         violations.append(
    #                             Violation(
    #                                 rule="16",
    #                                 code=code,
    #                                 message=f"{player} is within six months of the recorded {label} date.",
    #                                 team=origin,
    #                                 player=player,
    #                             )
    #                         )
    #
    #                 # Rules 17, 19, 22, 23 alter trade salary rather than always
    #                 # making a player categorically untradeable. Special salary
    #                 # fields are applied in salary matching; warn when the caller
    #                 # identifies a special case but omits the needed value.
    #                 if _as_bool(ov.get("poison_pill")) and ov.get("incoming_trade_salary") in (None, ""):
    #                     undetermined.append(
    #                         f"Rule 17: poison_pill=True for {origin} / {player}, but "
    #                         "incoming_trade_salary override is missing."
    #                     )
    #                 if pdata.guarantee_deadline and ov.get("outgoing_trade_salary") in (None, ""):
    #                     undetermined.append(
    #                         f"Rule 22: {origin} / {player} has guarantee-deadline data; "
    #                         "headline salary is being used because outgoing_trade_salary was not overridden."
    #                     )
    #                 if _as_bool(ov.get("has_trade_kicker")) and ov.get("incoming_trade_salary") in (None, ""):
    #                     undetermined.append(
    #                         f"Rule 23: has_trade_kicker=True for {origin} / {player}, but "
    #                         "incoming_trade_salary/apron_salary was not fully overridden."
    #                     )
    #
    #                 # Rule 18: CURRENT sign-and-trade event. Historical contract
    #                 # type alone does not trigger this branch.
    #                 if _as_bool(ov.get("being_signed_and_traded")):
    #                     self._validate_sign_and_trade_player(
    #                         origin,
    #                         destination,
    #                         player,
    #                         ov,
    #                         violations,
    #                         undetermined,
    #                     )
    #
    #                 # Rule 20: expiring contracts in the postseason trade window.
    #                 if self.rules.postseason_trade_window_open:
    #                     expiring = pdata.contract_end == self.rules.season_start
    #                     expiring = expiring or _as_bool(ov.get("could_expire_current_season"))
    #                     if expiring:
    #                         violations.append(
    #                             Violation(
    #                                 rule="20",
    #                                 code="POSTSEASON_EXPIRING_CONTRACT",
    #                                 message=(
    #                                     f"{player}'s contract expires, or may expire, with the current "
    #                                     "season and cannot be traded in the configured postseason window "
    #                                     "before the contract ends."
    #                                 ),
    #                                 team=origin,
    #                                 player=player,
    #                             )
    #                         )
    #
    #                 # Rule 21: reacquiring a traded-and-waived player.
    #                 blocked = ov.get("reacquire_blocked_until_by_team") or {}
    #                 if isinstance(blocked, Mapping) and destination in blocked:
    #                     until = _as_date(blocked[destination])
    #                     if until is not None and self.trade_date < until:
    #                         violations.append(
    #                             Violation(
    #                                 rule="21",
    #                                 code="REACQUISITION_WAITING_PERIOD",
    #                                 message=(
    #                                     f"{destination} cannot reacquire {player} until "
    #                                     f"{until.isoformat()}."
    #                                 ),
    #                                 team=destination,
    #                                 player=player,
    #                             )
    #                         )
    #
    #     def _validate_sign_and_trade_player(
    #         self,
    #         origin: str,
    #         destination: str,
    #         player: str,
    #         ov: Mapping[str, Any],
    #         violations: List[Violation],
    #         undetermined: List[str],
    #     ) -> None:
    #         # Rule 18 requirements directly listed in the supplied reference.
    #         bool_requirements = (
    #             ("sat_own_free_agent", "SIGN_AND_TRADE_NOT_OWN_FA", "player is not marked as the signing team's own free agent"),
    #             (
    #                 "sat_finished_prior_season_on_roster",
    #                 "SIGN_AND_TRADE_NOT_PRIOR_ROSTER",
    #                 "player is not marked as having finished the prior season on the signing team's roster",
    #             ),
    #             (
    #                 "sat_first_year_fully_guaranteed",
    #                 "SIGN_AND_TRADE_FIRST_YEAR_NOT_GUARANTEED",
    #                 "first season is not marked fully guaranteed",
    #             ),
    #         )
    #         for field_name, code, reason in bool_requirements:
    #             if field_name not in ov:
    #                 undetermined.append(
    #                     f"Rule 18: {field_name} missing for sign-and-trade player {origin} / {player}."
    #                 )
    #             elif not _as_bool(ov.get(field_name)):
    #                 violations.append(
    #                     Violation(
    #                         rule="18",
    #                         code=code,
    #                         message=f"Invalid sign-and-trade for {player}: {reason}.",
    #                         team=origin,
    #                         player=player,
    #                     )
    #                 )
    #
    #         excl = _as_int(ov.get("sat_contract_years_excluding_option"))
    #         incl = _as_int(ov.get("sat_contract_years_including_option"))
    #         if excl is None:
    #             undetermined.append(
    #                 f"Rule 18: sat_contract_years_excluding_option missing for {origin} / {player}."
    #             )
    #         elif excl < 3:
    #             violations.append(
    #                 Violation(
    #                     rule="18",
    #                     code="SIGN_AND_TRADE_TOO_SHORT",
    #                     message=f"{player}'s sign-and-trade contract has fewer than three seasons excluding an option year.",
    #                     team=origin,
    #                     player=player,
    #                 )
    #             )
    #         if incl is None:
    #             undetermined.append(
    #                 f"Rule 18: sat_contract_years_including_option missing for {origin} / {player}."
    #             )
    #         elif incl > 4:
    #             violations.append(
    #                 Violation(
    #                     rule="18",
    #                     code="SIGN_AND_TRADE_TOO_LONG",
    #                     message=f"{player}'s sign-and-trade contract exceeds four seasons including an option year.",
    #                     team=origin,
    #                     player=player,
    #                 )
    #             )
    #
    #         if self.rules.regular_season_start is None:
    #             undetermined.append(
    #                 f"Rule 18: regular_season_start is not configured for sign-and-trade {origin} / {player}."
    #             )
    #         elif self.trade_date >= self.rules.regular_season_start:
    #             violations.append(
    #                 Violation(
    #                     rule="18",
    #                     code="SIGN_AND_TRADE_AFTER_SEASON_START",
    #                     message=(
    #                         f"{player}'s sign-and-trade is dated {self.trade_date.isoformat()}, "
    #                         f"not before configured regular-season start {self.rules.regular_season_start.isoformat()}."
    #                     ),
    #                     team=origin,
    #                     player=player,
    #                 )
    #             )
    #
    #     # ------------------------- Rules 1-6, 17, 19, 22, 23, 31 -------------
    #
    #     def _validate_each_team_salary(
    #         self,
    #         flows: Mapping[str, TeamFlow],
    #         violations: List[Violation],
    #         undetermined: List[str],
    #         mechanisms: Dict[str, str],
    #         team_salary_detail: Dict[str, Dict[str, Money]],
    #     ) -> None:
    #         for team, flow in flows.items():
    #             tdata = self._team_data(team)
    #             if tdata is None:
    #                 violations.append(
    #                     Violation(
    #                         rule="31",
    #                         code="TEAM_CAP_DATA_MISSING",
    #                         message=f"No {self.rules.season_start} team cap/apron row found for {team}.",
    #                         team=team,
    #                     )
    #                 )
    #                 continue
    #
    #             missing_player = False
    #             outgoing_trade_salary = 0.0
    #             incoming_trade_salary = 0.0
    #             outgoing_apron_salary = 0.0
    #             incoming_apron_salary = 0.0
    #
    #             for player, _destination in flow.outgoing_players:
    #                 if self._player_data(team, player) is None:
    #                     missing_player = True
    #                     continue
    #                 outgoing_trade_salary += self._outgoing_trade_salary(team, player)
    #                 outgoing_apron_salary += self._apron_salary(team, player)
    #
    #             for player, origin in flow.incoming_players:
    #                 if self._player_data(origin, player) is None:
    #                     missing_player = True
    #                     continue
    #                 incoming_trade_salary += self._incoming_trade_salary(origin, player)
    #                 incoming_apron_salary += self._apron_salary(origin, player)
    #
    #             if missing_player:
    #                 undetermined.append(
    #                     f"Rules 1-6/31: salary matching skipped for {team} because at least one player row is missing."
    #                 )
    #                 continue
    #
    #             # The workbook's first_apron_allocations is the closest supplied
    #             # field to the current Apron Team Salary used by these rules.
    #             pre_apron_salary = tdata.first_apron_allocations
    #             post_apron_salary = pre_apron_salary - outgoing_apron_salary + incoming_apron_salary
    #
    #             first_apron = tdata.first_apron or self.rules.first_apron
    #             second_apron = tdata.second_apron or self.rules.second_apron
    #             cap_space = max(0.0, tdata.cap_space)
    #
    #             team_salary_detail[team] = {
    #                 "outgoing_trade_salary": outgoing_trade_salary,
    #                 "incoming_trade_salary": incoming_trade_salary,
    #                 "outgoing_apron_salary": outgoing_apron_salary,
    #                 "incoming_apron_salary": incoming_apron_salary,
    #                 "pretrade_apron_salary": pre_apron_salary,
    #                 "posttrade_apron_salary": post_apron_salary,
    #                 "cap_space": cap_space,
    #                 "first_apron": first_apron,
    #                 "second_apron": second_apron,
    #             }
    #
    #             # If team only sends salary and receives none, salary matching does
    #             # not prohibit the trade from this team's side.
    #             if incoming_trade_salary <= 0:
    #                 mechanisms[team] = "no_incoming_salary"
    #                 continue
    #
    #             receiving_sat = any(
    #                 _as_bool(self._player_override(origin, player).get("being_signed_and_traded"))
    #                 for player, origin in flow.incoming_players
    #             )
    #             explicit_first_apron_hard_cap = _as_bool(
    #                 self._team_override(team).get("hard_cap_first_apron")
    #             )
    #
    #             # Rule 18/5: receiving sign-and-trade player hard-caps at First Apron.
    #             if (receiving_sat or explicit_first_apron_hard_cap) and post_apron_salary > first_apron:
    #                 violations.append(
    #                     Violation(
    #                         rule="5/18",
    #                         code="FIRST_APRON_HARD_CAP_EXCEEDED",
    #                         message=(
    #                             f"{team} would have Apron Team Salary ${post_apron_salary:,.0f}, "
    #                             f"above First Apron ${first_apron:,.0f} after a First-Apron hard-cap trigger."
    #                         ),
    #                         team=team,
    #                     )
    #                 )
    #                 continue
    #
    #             mechanism = self._find_salary_matching_mechanism(
    #                 team=team,
    #                 flow=flow,
    #                 outgoing_salary=outgoing_trade_salary,
    #                 incoming_salary=incoming_trade_salary,
    #                 post_apron_salary=post_apron_salary,
    #                 first_apron=first_apron,
    #                 second_apron=second_apron,
    #                 cap_space=cap_space,
    #             )
    #
    #             if mechanism is None:
    #                 violations.append(
    #                     Violation(
    #                         rule="1-6/31",
    #                         code="SALARY_MATCHING_FAILED",
    #                         message=(
    #                             f"{team} cannot legally match incoming salary ${incoming_trade_salary:,.0f} "
    #                             f"against outgoing salary ${outgoing_trade_salary:,.0f} using the "
    #                             "cap-room, Standard TPE, Aggregated Standard TPE, Expanded TPE, or "
    #                             "available existing-TPE checks implemented from the supplied reference."
    #                         ),
    #                         team=team,
    #                     )
    #                 )
    #             else:
    #                 mechanisms[team] = mechanism
    #
    #     def _find_salary_matching_mechanism(
    #         self,
    #         *,
    #         team: str,
    #         flow: TeamFlow,
    #         outgoing_salary: Money,
    #         incoming_salary: Money,
    #         post_apron_salary: Money,
    #         first_apron: Money,
    #         second_apron: Money,
    #         cap_space: Money,
    #     ) -> Optional[str]:
    #         n_outgoing = len(flow.outgoing_players)
    #         cushion = self.rules.matching_cushion if post_apron_salary <= first_apron else 0.0
    #
    #         # Rule 1: team using cap space. Net incoming salary may exceed outgoing
    #         # by available room plus the listed $250K cushion.
    #         if cap_space > 0:
    #             cap_room_max = outgoing_salary + cap_space + self.rules.matching_cushion
    #             if incoming_salary <= cap_room_max:
    #                 return "cap_room"
    #
    #         # Rules 2-3 and 6: Standard matching. Multiple outgoing players may be
    #         # aggregated unless doing so leaves Apron Team Salary above Second Apron.
    #         aggregation_permitted = not (n_outgoing > 1 and post_apron_salary > second_apron)
    #         if aggregation_permitted:
    #             standard_max = outgoing_salary + cushion
    #             if incoming_salary <= standard_max:
    #                 return "aggregated_standard_tpe" if n_outgoing > 1 else "standard_tpe"
    #         elif self._can_match_without_aggregating_outgoing(team, flow):
    #             # A Second-Apron team may still have multiple players in the same
    #             # overall trade so long as the prohibited aggregation of outgoing
    #             # salaries is not needed. Treat each outgoing salary as its own
    #             # matching bucket; because the team is above the First Apron, no
    #             # $250K cushion is available in these buckets.
    #             return "second_apron_nonaggregated_standard"
    #
    #         # Rule 4: Expanded TPE tiers, available only if the resulting team can
    #         # obey the First-Apron hard cap that the mechanism itself triggers.
    #         if post_apron_salary <= first_apron:
    #             expanded_max = self._expanded_matching_max(outgoing_salary)
    #             if incoming_salary <= expanded_max:
    #                 return "expanded_tpe"
    #
    #         # Rule 31 mentions use of an existing TPE. Existing exceptions cannot
    #         # be inferred from the supplied workbook today (all sections are empty),
    #         # so callers can provide them via team_overrides. We use this only for
    #         # salary not already matched through outgoing salary. This conservative
    #         # implementation requires all incoming salary to fit a single TPE.
    #         for tpe in self._active_trade_exceptions(team):
    #             if tpe.amount >= incoming_salary:
    #                 if tpe.hard_caps_first_apron and post_apron_salary > first_apron:
    #                     continue
    #                 return "existing_tpe"
    #
    #         return None
    #
    #     def _can_match_without_aggregating_outgoing(
    #         self,
    #         team: str,
    #         flow: TeamFlow,
    #     ) -> bool:
    #         """
    #         Rule 6 helper for a team left above the Second Apron.
    #
    #         Outgoing salaries cannot be combined. We therefore ask whether every
    #         incoming player can be assigned wholly to an individual outgoing
    #         player's salary bucket, allowing more than one incoming player in a
    #         bucket when their combined salary fits that one outgoing salary.
    #
    #         This is intentionally different from summing all outgoing salaries.
    #         """
    #         capacities = sorted(
    #             (
    #                 self._outgoing_trade_salary(team, player)
    #                 for player, _destination in flow.outgoing_players
    #             ),
    #             reverse=True,
    #         )
    #         incoming = sorted(
    #             (
    #                 self._incoming_trade_salary(origin, player)
    #                 for player, origin in flow.incoming_players
    #             ),
    #             reverse=True,
    #         )
    #
    #         if not incoming:
    #             return True
    #         if not capacities or incoming[0] > capacities[0]:
    #             return False
    #
    #         # Small DFS/bin-packing search. Trade generators normally contain only
    #         # a handful of players per team, and symmetry pruning keeps this cheap.
    #         remaining = list(capacities)
    #
    #         def place(index: int) -> bool:
    #             if index >= len(incoming):
    #                 return True
    #
    #             salary = incoming[index]
    #             tried_remaining: set[float] = set()
    #             for i, capacity in enumerate(remaining):
    #                 if capacity < salary:
    #                     continue
    #                 rounded_capacity = round(capacity, 6)
    #                 if rounded_capacity in tried_remaining:
    #                     continue
    #                 tried_remaining.add(rounded_capacity)
    #
    #                 remaining[i] -= salary
    #                 if place(index + 1):
    #                     return True
    #                 remaining[i] += salary
    #
    #             return False
    #
    #         return place(0)
    #
    #     def _expanded_matching_max(self, outgoing_salary: Money) -> Money:
    #         if outgoing_salary <= self.rules.expanded_low_threshold:
    #             return 2.0 * outgoing_salary + self.rules.matching_cushion
    #         if outgoing_salary < self.rules.expanded_middle_threshold:
    #             return outgoing_salary + self.rules.expanded_middle_addon
    #         return 1.25 * outgoing_salary + self.rules.matching_cushion
    #
    #     def _active_trade_exceptions(self, team: str) -> List[TradeException]:
    #         exceptions = list(self.data.trade_exceptions.get(team, ()))
    #         override_entries = self._team_override(team).get("existing_tpes", ()) or ()
    #
    #         for entry in override_entries:
    #             if isinstance(entry, TradeException):
    #                 exceptions.append(entry)
    #                 continue
    #             if isinstance(entry, Mapping):
    #                 exceptions.append(
    #                     TradeException(
    #                         amount=_as_money(entry.get("amount")),
    #                         expires=_as_date(entry.get("expires")),
    #                         hard_caps_first_apron=_as_bool(entry.get("hard_caps_first_apron")),
    #                     )
    #                 )
    #
    #         return [
    #             tpe
    #             for tpe in exceptions
    #             if tpe.amount > 0 and (tpe.expires is None or self.trade_date <= tpe.expires)
    #         ]
    #
    #     # ------------------------- Rule 24 cash -------------------------------
    #
    #     def _validate_cash(
    #         self,
    #         flows: Mapping[str, TeamFlow],
    #         violations: List[Violation],
    #         undetermined: List[str],
    #         team_salary_detail: Mapping[str, Mapping[str, Money]],
    #     ) -> None:
    #         for team, flow in flows.items():
    #             if flow.cash_sent <= 0 and flow.cash_received <= 0:
    #                 continue
    #
    #             detail = team_salary_detail.get(team)
    #             if flow.cash_sent > 0 and detail is not None:
    #                 if detail["posttrade_apron_salary"] > detail["second_apron"]:
    #                     violations.append(
    #                         Violation(
    #                             rule="6/24",
    #                             code="SECOND_APRON_CASH_PROHIBITED",
    #                             message=f"{team} cannot send cash while the transaction leaves it above the Second Apron.",
    #                             team=team,
    #                         )
    #                     )
    #
    #             # A single trade above the annual limit is always impossible even
    #             # without YTD data.
    #             if flow.cash_sent > self.rules.annual_cash_limit:
    #                 violations.append(
    #                     Violation(
    #                         rule="24",
    #                         code="CASH_SENT_SINGLE_TRADE_OVER_LIMIT",
    #                         message=(
    #                             f"{team} sends ${flow.cash_sent:,.0f}, above the approximate "
    #                             f"2026-27 annual cash-sent limit ${self.rules.annual_cash_limit:,.0f}."
    #                         ),
    #                         team=team,
    #                     )
    #                 )
    #             if flow.cash_received > self.rules.annual_cash_limit:
    #                 violations.append(
    #                     Violation(
    #                         rule="24",
    #                         code="CASH_RECEIVED_SINGLE_TRADE_OVER_LIMIT",
    #                         message=(
    #                             f"{team} receives ${flow.cash_received:,.0f}, above the approximate "
    #                             f"2026-27 annual cash-received limit ${self.rules.annual_cash_limit:,.0f}."
    #                         ),
    #                         team=team,
    #                     )
    #                 )
    #
    #             tov = self._team_override(team)
    #             sent_ytd = tov.get("cash_sent_ytd")
    #             received_ytd = tov.get("cash_received_ytd")
    #
    #             if flow.cash_sent > 0:
    #                 if sent_ytd is None:
    #                     undetermined.append(
    #                         f"Rule 24: cash_sent_ytd missing for {team}; annual aggregate cash-sent limit cannot be fully checked."
    #                     )
    #                 elif _as_money(sent_ytd) + flow.cash_sent > self.rules.annual_cash_limit:
    #                     violations.append(
    #                         Violation(
    #                             rule="24",
    #                             code="CASH_SENT_ANNUAL_LIMIT",
    #                             message=f"{team} would exceed the annual cash-sent limit.",
    #                             team=team,
    #                         )
    #                     )
    #
    #             if flow.cash_received > 0:
    #                 if received_ytd is None:
    #                     undetermined.append(
    #                         f"Rule 24: cash_received_ytd missing for {team}; annual aggregate cash-received limit cannot be fully checked."
    #                     )
    #                 elif _as_money(received_ytd) + flow.cash_received > self.rules.annual_cash_limit:
    #                     violations.append(
    #                         Violation(
    #                             rule="24",
    #                             code="CASH_RECEIVED_ANNUAL_LIMIT",
    #                             message=f"{team} would exceed the annual cash-received limit.",
    #                             team=team,
    #                         )
    #                     )
    #
    #     # ------------------------- Rules 25-29 draft picks --------------------
    #
    #     def _validate_draft_assets(
    #         self,
    #         flows: Mapping[str, TeamFlow],
    #         violations: List[Violation],
    #         undetermined: List[str],
    #     ) -> None:
    #         if not any(flow.outgoing_picks or flow.incoming_picks for flow in flows.values()):
    #             return
    #
    #         max_year = self.rules.first_future_draft_year + 6
    #
    #         for team, flow in flows.items():
    #             frozen = {
    #                 int(year)
    #                 for year in (self._team_override(team).get("frozen_first_round_years", ()) or ())
    #             }
    #
    #             for pick in flow.outgoing_picks:
    #                 year = _as_int(pick.get("year"))
    #                 rnd = _as_int(pick.get("round"))
    #                 if year is None or rnd is None:
    #                     violations.append(
    #                         Violation(
    #                             rule="25-29",
    #                             code="PICK_YEAR_OR_ROUND_MISSING",
    #                             message="Outgoing pick must include integer year and round.",
    #                             team=team,
    #                         )
    #                     )
    #                     continue
    #
    #                 # Rule 25.
    #                 if year > max_year:
    #                     violations.append(
    #                         Violation(
    #                             rule="25",
    #                             code="PICK_BEYOND_SEVEN_DRAFTS",
    #                             message=(
    #                                 f"{team} attempts to trade a {year} pick, beyond the configured "
    #                                 f"seven-future-draft horizon ending in {max_year}."
    #                             ),
    #                             team=team,
    #                         )
    #                     )
    #
    #                 # Rule 29.
    #                 if rnd == 1 and year in frozen:
    #                     violations.append(
    #                         Violation(
    #                             rule="29",
    #                             code="FROZEN_FIRST_ROUND_PICK",
    #                             message=f"{team}'s {year} first-round pick is marked frozen/untradeable.",
    #                             team=team,
    #                         )
    #                     )
    #
    #                 # Rule 27: protected picks require possible conveyance years.
    #                 if rnd == 1 and _as_bool(pick.get("is_protected")):
    #                     possible = pick.get("possible_conveyance_years")
    #                     if not possible:
    #                         undetermined.append(
    #                             f"Rule 27: protected first-round pick from {team} in {year} has no "
    #                             "possible_conveyance_years/obligation-tree data; full Stepien validation is not possible."
    #                         )
    #
    #         self._validate_stepien(flows, violations, undetermined)
    #
    #     def _guaranteed_first_round_counts(self, team: str) -> Dict[int, int]:
    #         counts = {
    #             year: 0
    #             for year in range(
    #                 self.rules.first_future_draft_year,
    #                 self.rules.first_future_draft_year + 7,
    #             )
    #         }
    #
    #         for asset in self.data.draft_assets.get(team, ()): 
    #             if asset.round != 1 or asset.draft_year not in counts:
    #                 continue
    #
    #             # An unencumbered own pick or an unconditional incoming first is a
    #             # first-round selection the team possesses. An outbound swap also
    #             # leaves the team with a first-round pick, consistent with Rule 28.
    #             if asset.direction == "OWN":
    #                 counts[asset.draft_year] += 1
    #             elif asset.direction == "IN" and not asset.is_conditional and not asset.is_protected:
    #                 counts[asset.draft_year] += 1
    #             elif asset.direction == "OUT" and asset.is_swap:
    #                 counts[asset.draft_year] += 1
    #
    #         return counts
    #
    #     @staticmethod
    #     def _stepien_bad_pairs(counts: Mapping[int, int]) -> set[Tuple[int, int]]:
    #         years = sorted(counts)
    #         bad: set[Tuple[int, int]] = set()
    #         for a, b in zip(years, years[1:]):
    #             if b == a + 1 and counts[a] <= 0 and counts[b] <= 0:
    #                 bad.add((a, b))
    #         return bad
    #
    #     def _validate_stepien(
    #         self,
    #         flows: Mapping[str, TeamFlow],
    #         violations: List[Violation],
    #         undetermined: List[str],
    #     ) -> None:
    #         for team, flow in flows.items():
    #             first_round_moves = [
    #                 pick
    #                 for pick in flow.outgoing_picks
    #                 if _as_int(pick.get("round")) == 1 and not _as_bool(pick.get("is_swap"))
    #             ]
    #             incoming_firsts = [
    #                 pick
    #                 for pick in flow.incoming_picks
    #                 if _as_int(pick.get("round")) == 1 and not _as_bool(pick.get("is_swap"))
    #             ]
    #             if not first_round_moves and not incoming_firsts:
    #                 continue
    #
    #             base = self._guaranteed_first_round_counts(team)
    #             pre_bad = self._stepien_bad_pairs(base)
    #
    #             # Add clearly guaranteed incoming first-round picks.
    #             for pick in incoming_firsts:
    #                 year = _as_int(pick.get("year"))
    #                 if year not in base:
    #                     continue
    #                 if not _as_bool(pick.get("is_protected")) and not _as_bool(pick.get("is_conditional")):
    #                     base[year] += 1
    #
    #             deterministic_out: List[int] = []
    #             protected_options: List[List[Optional[int]]] = []
    #
    #             for pick in first_round_moves:
    #                 year = _as_int(pick.get("year"))
    #                 if year is None or year not in base:
    #                     continue
    #
    #                 if _as_bool(pick.get("is_protected")):
    #                     possible_raw = pick.get("possible_conveyance_years") or ()
    #                     possible = [
    #                         int(y)
    #                         for y in possible_raw
    #                         if int(y) in base
    #                     ]
    #                     if possible:
    #                         # Include None for a non-conveyance/extinguishment scenario.
    #                         protected_options.append([None] + sorted(set(possible)))
    #                     continue
    #
    #                 deterministic_out.append(year)
    #
    #             # If a matching exact asset is supplied and the team owns multiple
    #             # firsts in a year, this one decrement correctly preserves another.
    #             baseline_after_deterministic = dict(base)
    #             for year in deterministic_out:
    #                 baseline_after_deterministic[year] = max(0, baseline_after_deterministic[year] - 1)
    #
    #             scenarios: Iterable[Tuple[Optional[int], ...]]
    #             if protected_options:
    #                 # Avoid pathological explosion in future expanded generators.
    #                 scenario_count = 1
    #                 for options in protected_options:
    #                     scenario_count *= len(options)
    #                 if scenario_count > 4096:
    #                     undetermined.append(
    #                         f"Rule 27: protected-pick Stepien scenario count for {team} exceeds 4096; "
    #                         "full obligation-tree enumeration was skipped."
    #                     )
    #                     scenarios = [tuple()]
    #                 else:
    #                     scenarios = cartesian_product(*protected_options)
    #             else:
    #                 scenarios = [tuple()]
    #
    #             for scenario in scenarios:
    #                 counts = dict(baseline_after_deterministic)
    #                 for conveyed_year in scenario:
    #                     if conveyed_year is not None:
    #                         counts[conveyed_year] = max(0, counts[conveyed_year] - 1)
    #
    #                 post_bad = self._stepien_bad_pairs(counts)
    #                 new_bad = post_bad - pre_bad
    #                 if new_bad:
    #                     pair = sorted(new_bad)[0]
    #                     violations.append(
    #                         Violation(
    #                             rule="26-27",
    #                             code="STEPIEN_RULE",
    #                             message=(
    #                                 f"{team}'s outgoing first-round pick package can create consecutive "
    #                                 f"future drafts {pair[0]} and {pair[1]} with no first-round selection."
    #                             ),
    #                             team=team,
    #                         )
    #                     )
    #                     break
    #
    #
    # # ===========================================================================
    # # Only public function: two arguments, no API calls, no separate modules.
    # # ===========================================================================
    #
    # def generate_legal_two_team_trades(
    #     workbook_path: str | Path,
    #     user_preferences_string: str,
    # ) -> Iterator[Dict[str, Any]]:
    #     """Generate all two-team, player-for-player trades passing known CBA checks.
    #
    #     Parameters
    #     ----------
    #     workbook_path
    #         The supplied NBA engine .xlsx file.
    #     user_preferences_string
    #         Preferred: a JSON (or Python-literal) string with these fields:
    #
    #             {
    #               "user_team": "New York Knicks",
    #               "stage_1_players": {"New York Knicks": [...], "Atlanta Hawks": [...]},
    #               "stage_2_targets": {"Atlanta Hawks": [...]},
    #               "trade_date": "2026-10-08",  # optional; default: local today
    #               "player_overrides": {"Team|Player": {"signed_date": "..."}},
    #               "team_overrides": {"Team": {"existing_tpes": [...]}},
    #               "season_rules": {"trade_deadline": "2027-02-04"}
    #             }
    #
    #         Also accepts the upstream raw output of two consecutive Python dicts,
    #         IF prefixed by a line such as "Team Name: New York Knicks".
    #         The first dict is Stage 1 (allowable players). The second is Stage 2
    #         (desired targets). All received players must be in Stage 2.
    #
    #     Returns
    #     -------
    #     Iterator of trade dictionaries in the previous 'trade_type'/'teams'/'moves'
    #     format, with a 'cba_validation' diagnostics field. Consume with a for-loop.
    #
    #     Generation policy
    #     -----------------
    #     * Exactly two teams, always including the user's team.
    #     * Each team sends >=1 player and receives >=1 player.
    #     * Only Stage 1 allowable players may leave their original teams.
    #     * All user-team acquisitions must be Stage 2 targets.
    #     * Only 2026-27 workbook contract rows count for the current season.
    #     * No draft picks or cash are ADDED by this generator.
    #     * Player pre-eligibility, permissive salary ceilings, full CBA validation.
    #     * No basketball-utility optimization or certification based on missing data.
    #     """
    #     import ast
    #     import json
    #     import re
    #     from bisect import bisect_right
    #     from dataclasses import fields, replace
    #     from itertools import combinations
    #
    #     if not isinstance(user_preferences_string, str) or not user_preferences_string.strip():
    #         raise ValueError("user_preferences_string must be a nonempty string.")
    #
    #     raw = user_preferences_string.strip()
    #     payload = None
    #     try:
    #         payload = json.loads(raw)
    #     except (ValueError, TypeError):
    #         try:
    #             literal = ast.literal_eval(raw)
    #             if isinstance(literal, dict):
    #                 payload = literal
    #         except (ValueError, SyntaxError):
    #             pass
    #
    #     if isinstance(payload, dict):
    #         user_team = payload.get("user_team") or payload.get("team_name")
    #         stage1 = payload.get("stage_1_players", payload.get("tradeable_players"))
    #         stage2 = payload.get("stage_2_targets", payload.get("wanted_players"))
    #         trade_date_raw = payload.get("trade_date")
    #         rules_config = payload.get("season_rules") or {}
    #         player_overrides_raw = payload.get("player_overrides") or {}
    #         team_overrides = payload.get("team_overrides") or {}
    #     else:
    #         # Parse the two UNASSIGNED dictionary expressions from GPT's Stage1/2
    #         # response. Team name must be provided in the text prefix.
    #         name_match = re.search(
    #             r"(?im)^\s*(?:Team Name|User Team|user_team)\s*:\s*[\"']?(.+?)[\"']?\s*$",
    #             raw,
    #         )
    #         if not name_match:
    #             raise ValueError(
    #                 "The preference string must identify the user team. Supply "
    #                 "JSON with user_team, stage_1_players, stage_2_targets, or "
    #                 "prefix the two raw Python dictionaries with 'Team Name: ...'."
    #             )
    #         user_team = name_match.group(1).strip().strip("\"'")
    #         start = raw.find("{")
    #         if start < 0:
    #             raise ValueError("Cannot find the two Stage 1/Stage 2 dictionaries.")
    #         try:
    #             tree = ast.parse(raw[start:], mode="exec")
    #             dicts = [
    #                 ast.literal_eval(node.value)
    #                 for node in tree.body
    #                 if isinstance(node, ast.Expr)
    #                 and isinstance(node.value, ast.Dict)
    #             ]
    #         except (SyntaxError, ValueError) as exc:
    #             raise ValueError("Could not parse the two Stage 1/Stage 2 dictionaries.") from exc
    #         if len(dicts) != 2:
    #             raise ValueError("Expected exactly TWO upstream Python dictionaries.")
    #         stage1, stage2 = dicts
    #         trade_date_raw = None
    #         rules_config = {}
    #         player_overrides_raw = {}
    #         team_overrides = {}
    #
    #     if not isinstance(user_team, str) or not user_team.strip():
    #         raise ValueError("Include user_team in the preference string.")
    #     user_team = user_team.strip()
    #     if not isinstance(stage1, dict) or not isinstance(stage2, dict):
    #         raise ValueError("Stage 1 and Stage 2 must both be dictionaries of team -> player list.")
    #     if user_team not in stage1:
    #         raise ValueError(f"User team {user_team!r} is missing from Stage 1.")
    #
    #     def normalize_player_dict(data: dict, label: str) -> Dict[str, Tuple[str, ...]]:
    #         normalized = {}
    #         for team, names in data.items():
    #             if not isinstance(team, str) or not isinstance(names, (list, tuple)):
    #                 raise ValueError(f"{label} must map team names to lists of player names.")
    #             if any(not isinstance(p, str) or not p.strip() for p in names):
    #                 raise ValueError(f"{label} contains an invalid player name for {team}.")
    #             normalized[team.strip()] = tuple(sorted(set(p.strip() for p in names)))
    #         return normalized
    #
    #     allowable = normalize_player_dict(stage1, "Stage 1")
    #     targets = normalize_player_dict(stage2, "Stage 2")
    #
    #     if user_team in targets:
    #         raise ValueError("Stage 2 must contain opposing-team targets only.")
    #     for team, names in targets.items():
    #         if team not in allowable:
    #             raise ValueError(f"Stage 2 includes {team!r}, which was removed in Stage 1.")
    #         missing = set(names) - set(allowable[team])
    #         if missing:
    #             raise ValueError(
    #                 f"Stage 2 contains excluded/non-allowable player(s) on {team}: "
    #                 f"{sorted(missing)!r}. Fix the upstream preference output."
    #             )
    #
    #     if trade_date_raw:
    #         execution_date = _as_date(trade_date_raw)
    #     else:
    #         execution_date = date.today()
    #
    #     if not isinstance(rules_config, dict):
    #         raise ValueError("season_rules must be a dictionary when supplied.")
    #     valid_rule_keys = {field.name for field in fields(SeasonRules)}
    #     unknown_rule_keys = set(rules_config) - valid_rule_keys
    #     if unknown_rule_keys:
    #         raise ValueError(f"Unrecognized season_rules keys: {sorted(unknown_rule_keys)!r}")
    #     rules_config = dict(rules_config)
    #     for key in ("trade_deadline", "regular_season_start"):
    #         if key in rules_config:
    #             rules_config[key] = _as_date(rules_config[key])
    #     rules = replace(SeasonRules(), **rules_config)
    #     if rules.season_start != 2026:
    #         raise ValueError(
    #             "This integrated ruleset has 2026-27 financial thresholds. "
    #             "Update the source rules before using it for a different season."
    #         )
    #
    #     if not isinstance(player_overrides_raw, dict) or not isinstance(team_overrides, dict):
    #         raise ValueError("player_overrides and team_overrides must be dictionaries.")
    #
    #     player_overrides = {}
    #     for key, overrides in player_overrides_raw.items():
    #         if not isinstance(overrides, dict):
    #             raise ValueError("Each player override value must be a dictionary.")
    #         if isinstance(key, (tuple, list)) and len(key) == 2:
    #             player_key = (str(key[0]), str(key[1]))
    #         elif isinstance(key, str) and "|" in key:
    #             player_key = tuple(part.strip() for part in key.split("|", 1))
    #         else:
    #             raise ValueError(
    #                 "Player override keys must be 'Team|Player' (or a 2-tuple "
    #                 "in a Python-literal preference dictionary)."
    #             )
    #         player_overrides[player_key] = overrides
    #
    #     validator = CBAValidator.from_xlsx(
    #         workbook_path,
    #         trade_date=execution_date,
    #         rules=rules,
    #         player_overrides=player_overrides,
    #         team_overrides=team_overrides,
    #     )
    #     data = validator.data
    #
    #     # Prevent stale/wrong rosters from silently becoming legal candidates.
    #     for team, names in allowable.items():
    #         if team not in data.teams:
    #             raise ValueError(f"Stage 1 team {team!r} is absent from current workbook team data.")
    #         for player in names:
    #             if (team, player) not in data.players:
    #                 raise ValueError(
    #                     f"Stage 1 player {player!r} has no 2026-27 contract record "
    #                     f"for {team!r}. Re-run upstream roster/preference collection."
    #                 )
    #
    #     # Only remove a player early if the supplied validator can ALREADY PROVE
    #     # a player-level violation from a single outgoing move. Restrictions on
    #     # aggregated or multi-player deals are still checked on full candidates.
    #     eligibility_cache: Dict[Tuple[str, str, str], bool] = {}
    #
    #     def individually_eligible(origin: str, name: str, destination: str) -> bool:
    #         key = (origin, name, destination)
    #         if key not in eligibility_cache:
    #             flow = TeamFlow(
    #                 outgoing_players=((name, destination),),
    #                 incoming_players=(),
    #                 cash_sent=0.0,
    #                 cash_received=0.0,
    #                 outgoing_picks=(),
    #                 incoming_picks=(),
    #             )
    #             violations: List[Violation] = []
    #             unknown: List[str] = []
    #             validator._validate_player_ownership_and_eligibility(
    #                 {origin: flow}, violations, unknown
    #             )
    #             eligibility_cache[key] = not violations
    #         return eligibility_cache[key]
    #
    #     # Financial pruning uses a DELIBERATELY PERMISSIVE upper bound. It can
    #     # admit candidates that the full validator rejects, but it must never
    #     # reject a trade that the original salary validator could have accepted.
    #     max_tpe_by_team: Dict[str, float] = {}
    #
    #     def permissive_incoming_ceiling(team: str, outgoing: float) -> float:
    #         tdata = data.teams[team]
    #         if team not in max_tpe_by_team:
    #             max_tpe_by_team[team] = max(
    #                 [0.0] + [tpe.amount for tpe in validator._active_trade_exceptions(team)]
    #             )
    #         return max(
    #             outgoing + max(0.0, tdata.cap_space) + rules.matching_cushion,
    #             outgoing + rules.matching_cushion,
    #             validator._expanded_matching_max(outgoing),
    #             max_tpe_by_team[team],
    #         )
    #
    #     def user_subsets(available_names: Tuple[str, ...]):
    #         # Outgoing and incoming salary can differ because of poison-pill,
    #         # kicker, BYC, or other optional salary override mechanisms.
    #         subsets = [(0.0, 0.0, ())]
    #         for player in available_names:
    #             outgoing = validator._outgoing_trade_salary(user_team, player)
    #             incoming = validator._incoming_trade_salary(user_team, player)
    #             subsets += [
    #                 (out_sum + outgoing, in_sum + incoming, names + (player,))
    #                 for out_sum, in_sum, names in subsets
    #             ]
    #         return sorted(subsets[1:], key=lambda part: part[0])
    #
    #     unknown_global = []
    #     if rules.trade_deadline is None:
    #         unknown_global.append(
    #             "Rule 30: no 2026-27 trade deadline date was provided; "
    #             "the trade-deadline check is undetermined."
    #         )
    #
    #     subset_cache = {}
    #
    #     def iter_results() -> Iterator[Dict[str, Any]]:
    #         for other_team in sorted(targets):
    #             if other_team == user_team or not targets[other_team]:
    #                 continue
    #
    #             user_names = tuple(
    #                 name for name in allowable[user_team]
    #                 if individually_eligible(user_team, name, other_team)
    #             )
    #             target_names = tuple(
    #                 name for name in targets[other_team]
    #                 if individually_eligible(other_team, name, user_team)
    #             )
    #             if not user_names or not target_names:
    #                 continue
    #
    #             # Reuse user subset sums for every target combination on this team.
    #             if user_names not in subset_cache:
    #                 subset_cache[user_names] = user_subsets(user_names)
    #             outbound_subsets = subset_cache[user_names]
    #             if not outbound_subsets:
    #                 continue
    #             outbound_salaries = [part[0] for part in outbound_subsets]
    #             no_special_user_salary = all(
    #                 abs(validator._incoming_trade_salary(user_team, name)
    #                     - validator._outgoing_trade_salary(user_team, name)) < 0.01
    #                 for name in user_names
    #             )
    #
    #             # Targets-only incoming subsets; no unwanted filler players.
    #             # combinations() does not construct the full target power set.
    #             for target_count in range(1, len(target_names) + 1):
    #                 for received in combinations(target_names, target_count):
    #                     incoming_to_user = sum(
    #                         validator._incoming_trade_salary(other_team, p)
    #                         for p in received
    #                     )
    #                     outgoing_from_other = sum(
    #                         validator._outgoing_trade_salary(other_team, p)
    #                         for p in received
    #                     )
    #
    #                     # Find the FIRST user outgoing subset that could possibly
    #                     # match the required incoming salary. Monotonic bound.
    #                     lo, hi = 0, len(outbound_subsets)
    #                     while lo < hi:
    #                         mid = (lo + hi) // 2
    #                         if permissive_incoming_ceiling(
    #                             user_team, outbound_salaries[mid]
    #                         ) < incoming_to_user - 0.01:
    #                             lo = mid + 1
    #                         else:
    #                             hi = mid
    #                     min_index = lo
    #
    #                     # An opponent must ALSO be able to receive the user's
    #                     # salary. When outbound and inbound salary match for
    #                     # user players, bisect out impossible large offers.
    #                     other_team_ceiling = permissive_incoming_ceiling(
    #                         other_team, outgoing_from_other
    #                     )
    #                     max_index = (
    #                         bisect_right(outbound_salaries, other_team_ceiling + 0.01)
    #                         if no_special_user_salary
    #                         else len(outbound_subsets)
    #                     )
    #                     if min_index >= max_index:
    #                         continue
    #
    #                     for idx in range(min_index, max_index):
    #                         outgoing_sum, opponent_incoming, sent = outbound_subsets[idx]
    #                         if opponent_incoming > other_team_ceiling + 0.01:
    #                             continue
    #
    #                         trade = {
    #                             "trade_type": "2_team",
    #                             "teams": (user_team, other_team),
    #                             "moves": (
    #                                 {
    #                                     "from_team": user_team,
    #                                     "to_team": other_team,
    #                                     "players": sent,
    #                                 },
    #                                 {
    #                                     "from_team": other_team,
    #                                     "to_team": user_team,
    #                                     "players": received,
    #                                 },
    #                             ),
    #                         }
    #                         result = validator.validate_trade(trade)
    #                         if not result.passes:
    #                             continue
    #
    #                         # Keep the validator's original financial diagnostics,
    #                         # and explicitly flag checks requiring missing facts.
    #                         unresolved = list(dict.fromkeys(
    #                             unknown_global + result.undetermined
    #                         ))
    #                         trade["cba_validation"] = {
    #                             "status": (
    #                                 "no_detected_violation_with_unresolved_checks"
    #                                 if unresolved
    #                                 else "no_detected_violation"
    #                             ),
    #                             "unresolved_checks": unresolved,
    #                             "salary_mechanisms": result.mechanisms,
    #                             "team_salary_detail": result.team_salary_detail,
    #                         }
    #                         yield trade
    #
    #     return iter_results()
    #
    #
    # # USAGE (upstream OpenAI parsing is performed in YOUR existing pipeline):
    # #
    # import json
    # preference_string = json.dumps({
    #      "user_team": "New York Knicks",
    #      "stage_1_players": stage_1_players,
    #      "stage_2_targets": stage_2_targets,
    #      "trade_date": "2026-10-08",
    #  })
    # #
    # trades = generate_legal_two_team_trades("nba_trade_engine_data.xlsx", preference_string)
    # for trade in trades:
    #      print(trade)
